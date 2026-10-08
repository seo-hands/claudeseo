#!/usr/bin/env python3
"""Shared helpers of the tz-travel skill (stage 2: meta tags + copywriter brief for one page of the stage-1 distribution).

Depends on the semantics-travel skill: settings, market profile (scripts/profiles/<code>.json, section "tz") and
decisions are read with its sp_common.  Site data comes from the nearest "## Дані сайту" (working folder, then its
parents); the direction settings and the journals come only from the working folder's CLAUDE.md.  Stage-1 files,
tz-config.json and all outputs live in the working folder (--workdir).  Nothing here calls DataForSEO.
"""
import glob, hashlib, json, os, re, sys, unicodedata
from urllib.parse import urlsplit

import openpyxl

SKILL_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SEMANTICS_SCRIPTS = os.path.join(os.path.dirname(SKILL_DIR), "semantics-travel", "scripts")
if not os.path.exists(os.path.join(SEMANTICS_SCRIPTS, "sp_common.py")):
    sys.exit(f"tz-travel потребує скіл semantics-travel: не знайдено {SEMANTICS_SCRIPTS}\\sp_common.py")
sys.path.insert(0, SEMANTICS_SCRIPTS)
import sp_common as sp  # noqa: E402

CONFIG_FILE = "tz-config.json"
DEFAULTS = {
    "competitors_limit": 12,        # pages to analyse
    "adjacent_keywords": 4,         # keywords next to the main one that form the competitor list
    "reserves_min": 3,
    "thr_core": 0.80, "thr_mid": 0.65,      # keyword similarity tiers (a hint, not a decision)
    "term_min_competitors": 3,      # semantic word: present on at least N competitor pages
    "term_max": 150,
    "topic_threshold": 0.65,        # clustering of competitor H2/H3
    "gap_min_competitors": 5,       # gap: topic on at least N competitors and absent on the own page
    "topic_text_share": 0.5,        # share of topic words that must occur in the own body text
    "cookie_wall_text_max": 500,
    "models_dir": "E:/Work/claudeseo/models",
    "brand": "",                    # brand suffix for the title; default: host of `site`
    "own_text_exclude": [],         # regexes of UI messages in the own page HTML that are not page text
    "browser_user_agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/141.0.0.0 Safari/537.36",
    "render_timeout_ms": 45000,
}
SHEET_KEYS, SHEET_PAGES, SHEET_COMP = "Ключі", "Розподіл по сторінках", "Конкуренти"


def norm(s):
    return re.sub(r"\s+", " ", (s or "").replace("\xa0", " ").replace("\u00ad", "")).strip()


def nosp(s):
    return len(re.sub(r"\s+", "", s or ""))


def host(u):
    return urlsplit(u).netloc.replace("www.", "")


def symbols(text):
    """pictographic symbols (✓ ✈ ☀ ✅ ➤ ...); ordinary punctuation, currency and maths signs are not counted"""
    return [ch for ch in text or "" if unicodedata.category(ch) in ("So", "Sk") or 0x1F000 <= ord(ch) <= 0x1FAFF]


def rjson(path):
    return json.load(open(path, encoding="utf-8"))


def wjson(path, data):
    json.dump(data, open(path, "w", encoding="utf-8"), ensure_ascii=False, indent=1)


def die(msg):
    sys.exit("tz-travel: " + msg)


def _sheet(wb, name):
    if name not in wb.sheetnames:
        die(f"в xlsx етапу 1 немає аркуша «{name}». Спершу запустіть semantics-travel (analyze.py).")
    rows = list(wb[name].iter_rows(values_only=True))
    head = [norm(str(h or "")) for h in rows[0]]
    return head, rows[1:]


def _col(head, *names):
    for n in names:
        for i, h in enumerate(head):
            if h.startswith(n):
                return i
    die(f"не знайдено колонку {names} серед {head}")


def page_path(u):
    return (urlsplit(u).path if "://" in u else u).rstrip("/") or "/"


def page_slug(S, path):
    base = S["page_base"].rstrip("/")
    if path == base:
        return S["slug"]
    if path.startswith(base + "/"):
        return path[len(base) + 1:].replace("/", "-")
    return "-".join(path.strip("/").split("/")[-2:])


def load_ctx(workdir, page):
    """Everything stage 1 knows about the page: keywords, landing type, neighbour pages and their main words."""
    try:
        S = sp.load_settings(workdir)
        P = sp.load_profile(S, workdir)
    except (sp.ConfigMissing, sp.ProfileMissing) as e:
        die(str(e))
    T = P.get("tz")
    if not T:
        die(f"у профілі ринку «{sp.profile_code(S)}» немає розділу \"tz\" (ліміти мета-тегів, головні слова сторінок, шаблони розбору). "
            "Скопіюйте розділ з profiles/de.json і перекладіть слова — див. references/rules.md скіла tz-travel.")
    cfg = dict(DEFAULTS)
    if os.path.exists(os.path.join(workdir, CONFIG_FILE)):
        cfg.update({k: v for k, v in rjson(os.path.join(workdir, CONFIG_FILE)).items() if not k.startswith("_")})
    cfg["brand"] = cfg["brand"] or host(S["site"])
    if not S.get("_direction_source"):
        die(f"у {os.path.join(os.path.abspath(workdir), 'CLAUDE.md')} немає секції «## semantics-travel»: для цієї папки етап 1 не налаштовано.")
    xlsx = os.path.join(workdir, f"semantics-{S['slug']}.xlsx")
    if not os.path.exists(xlsx):
        die(f"немає {xlsx}: етап 1 (semantics-travel) для цієї папки не зроблено.")
    wb = openpyxl.load_workbook(xlsx, read_only=True)
    path = page_path(page)

    head, rows = _sheet(wb, SHEET_PAGES)
    c_page, c_state, c_sum, c_type = _col(head, "сторінка"), _col(head, "стан"), _col(head, "сума"), _col(head, "домінуючий тип")
    pages = {}
    for r in rows:
        if r[c_page]:
            pages.setdefault(page_path(str(r[c_page])), {"state": norm(str(r[c_state] or "")).split(" ")[0], "volume": r[c_sum] or 0,
                                                         "type": norm(str(r[c_type] or "")).split(":")[0].strip("—- ") or None})
    if path not in pages:
        die(f"сторінки {path} немає в аркуші «{SHEET_PAGES}». Є: " + ", ".join(pages))

    head, rows = _sheet(wb, SHEET_KEYS)
    ck, ct, cv, cp = _col(head, "ключ"), _col(head, "переклад"), _col(head, "частотність"), _col(head, "рекомендована сторінка")
    cty, csh, cst, ccl = _col(head, "тип сторінки"), _col(head, "частка"), _col(head, "статус"), _col(head, "кластер")
    by_page = {}
    for r in rows:
        if r[ck] and r[cp]:
            by_page.setdefault(page_path(str(r[cp])), []).append({"keyword": r[ck], "translation": r[ct], "volume": r[cv] or 0, "type_top": r[cty],
                                                                  "share": r[csh], "status": r[cst], "cluster": norm(str(r[ccl] or "")).split(" ")[0]})
    keys = sorted(by_page.get(path, []), key=lambda k: -k["volume"])
    if not keys:
        die(f"на сторінку {path} етап 1 не призначив жодного ключа (батьківська сторінка матриць?). ТЗ без ключів не будується.")
    main = keys[0]["keyword"]
    clusters = {}      # stage-1 clusters assigned to the page: each one must get its own H2 or H3 in the brief
    for k in keys:     # keys are sorted by volume, so the first key of a cluster is its head keyword
        c = clusters.setdefault(k["cluster"], {"id": k["cluster"], "head": k["keyword"], "volume": k["volume"], "sum": 0, "keywords": []})
        c["sum"] += k["volume"]
        c["keywords"].append(k["keyword"])
    if path == S["page_base"].rstrip("/") and S.get("main_keyword") and any(k["keyword"] == S["main_keyword"] for k in keys):
        main = S["main_keyword"]

    head, rows = _sheet(wb, SHEET_COMP)
    c1, c2, c3 = _col(head, "ключ"), _col(head, "URL"), _col(head, "тип посадкової")
    c4 = _col(head, "title")
    url_type = {(r[c1], r[c2]): r[c3] for r in rows if r[c2]}
    serp_title = {r[c2]: r[c4] for r in rows if r[c2] and r[c4]}

    # neighbour pages: product pages give "head words" that are banned in the target's title and headings
    neighbours, regions = [], dict(S.get("regions") or {})
    for p, info in pages.items():
        if p == path:
            continue
        pk = sorted(by_page.get(p, []), key=lambda k: -k["volume"])
        last = p.rstrip("/").split("/")[-1]
        n = {"page": p, "state": info["state"], "volume": info["volume"], "main_keyword": pk[0]["keyword"] if pk else None, "kind": "other", "label": None, "match": None}
        fam = next((h for h in T["head_words"] if last == P.get("page_slugs", {}).get(h["family"]) or re.search(h["match"], last.replace("-", " "), re.I)), None)
        inf = next((t for t in T["info_topics"] if t["slug"] in last), None)
        if fam:
            n.update(kind="product", label=fam["label"], match=fam["match"])
        elif inf:
            n.update(kind="info", label=inf["label"], match=inf["match"])
        elif any(re.search(rx, last, re.I) for rx in regions.values()) or last in (S.get("lemma_regions") or "").split("|"):
            n["kind"] = "region"
        neighbours.append(n)
    banned = []
    for n in neighbours:
        if n["kind"] == "product" and not re.search(n["match"], main, re.I) and not re.search(n["match"], path.replace("-", " "), re.I):
            if n["label"] not in [b["label"] for b in banned]:
                banned.append({"label": n["label"], "match": n["match"], "page": n["page"]})
    slug = page_slug(S, path)
    return {"S": S, "P": P, "T": T, "cfg": cfg, "workdir": workdir, "path": path, "url": S["site"] + path, "slug": slug, "state": pages[path]["state"],
            "landing_type": pages[path]["type"], "keys": keys, "main": main, "clusters": clusters, "neighbours": neighbours, "banned": banned,
            "url_type": url_type, "serp_title": serp_title, "decisions": sp.load_decisions(workdir),
            "f": {"selection": os.path.join(workdir, f"competitors-{slug}.selection.json"), "competitors": os.path.join(workdir, f"competitors-{slug}.json"),
                  "embeddings": os.path.join(workdir, f"embeddings-{slug}.json"), "plan": os.path.join(workdir, f"tz-{slug}-plan.json"),
                  "draft": os.path.join(workdir, f"tz-{slug}-plan.draft.json"), "raw": os.path.join(workdir, "competitors-raw")}}


def load_serps(workdir):
    """keyword -> [(position, url, title)] from the stage-1 cache; no API calls"""
    out = {}
    for f in glob.glob(os.path.join(workdir, "serp-raw-regular", "*.json")):
        d = rjson(f)
        try:
            items = d["response"]["tasks"][0]["result"][0]["items"] or []
        except (KeyError, IndexError, TypeError):
            continue
        out[d["keyword"]] = [(i.get("rank_group"), i["url"], i.get("title")) for i in items if i.get("type") == "organic"]
    return out


def flag_keyword(ctx, kw):
    """label of a neighbour page's head word contained in the keyword (such a key never goes to title/H1)"""
    for b in ctx["banned"]:
        if re.search(b["match"], kw, re.I):
            return b["label"]
    return None


def cache_file(raw_dir, url, mode):
    u = urlsplit(url)
    return os.path.join(raw_dir, re.sub(r"\W+", "_", u.netloc + u.path)[:70] + "_" + hashlib.md5((url + mode).encode()).hexdigest()[:6] + ".json")


def exact_count(kw, text):
    return len(re.findall(r"(?<!\w)" + re.escape(kw.lower()) + r"(?!\w)", (text or "").lower()))


def args_page(ap):
    ap.add_argument("--workdir", default=".")
    ap.add_argument("--page", required=True, help="URL або шлях цільової сторінки з аркуша «Розподіл по сторінках»")
    return ap
