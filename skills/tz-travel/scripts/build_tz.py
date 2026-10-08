#!/usr/bin/env python3
"""Step d: meta-<slug>.xlsx and tz-copywriter-<slug>.docx from the analysis, the embeddings and the checked plan.

python build_tz.py --workdir . --page <URL> [--suffix .new] [--no-claude-md]

All checks of plan_check.py run first; if one fails the script stops with the explanation and writes nothing.
--suffix .new writes meta-<slug>.new.xlsx / tz-copywriter-<slug>.new.docx and leaves approved files untouched.
"""
import argparse, datetime, os, re, statistics, sys

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import plan_check as pc
import tz_common as tz

GAP_TAG, REQ_TAG = "прогалина за конкурентами", "обов'язковий за ключами"
HEAD_FILL, GAP_FILL, OK_FILL, GREY = (PatternFill("solid", fgColor=c) for c in ("1F3864", "FCE4D6", "E2EFDA", "EDEDED"))
num = lambda n: f"{n:,}".replace(",", " ")
chars = lambda c: f"{num(c[0])}–{num(c[1])}" if c else "—"
LEVEL = {"intro": "вступ", "listing": "лістинг"}


def sheet(wb, name, header, rows, widths, note=None):
    ws = wb.create_sheet(name)
    r0 = 1
    if note:
        ws.cell(1, 1, note).alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(header))
        ws.row_dimensions[1].height = 15 * (1 + len(note) // 150)
        r0 = 2
    for j, h in enumerate(header, 1):
        c = ws.cell(r0, j, h)
        c.font, c.fill, c.alignment = Font(bold=True, color="FFFFFF"), HEAD_FILL, Alignment(wrap_text=True, vertical="center")
    for i, row in enumerate(rows, r0 + 1):
        for j, v in enumerate(row, 1):
            ws.cell(i, j, v).alignment = Alignment(wrap_text=True, vertical="top")
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = ws.cell(r0 + 1, 1)
    ws.auto_filter.ref = f"A{r0}:{get_column_letter(len(header))}{r0 + len(rows)}"
    return ws, r0


def fill_row(ws, i, n, fill):
    for j in range(1, n + 1):
        ws.cell(i, j).fill = fill


class Build:
    def __init__(self, ctx, plan, C, E):
        self.ctx, self.plan, self.C, self.E = ctx, plan, C, E
        self.T, self.cfg, self.S = ctx["T"], ctx["cfg"], C["summary"]
        self.comps = [c for c in C["competitors"] if c.get("ok")]
        self.own, self.N, self.main = C["own"], len([c for c in C["competitors"] if c.get("ok")]), ctx["main"]
        self.topics = pc.topics(ctx, C, E)
        self.heads, self.faqx, self.textx = pc.exact_plan(ctx, plan)
        self.lo, self.hi = pc.sums(plan)
        self.site_host = tz.host(ctx["S"]["site"])
        self.sets = [dict(plan["meta"]["A"], name="Комплект A", status="впроваджувати", condition=plan["meta"]["A"].get("condition", "Впроваджувати зараз."))] + \
                    [dict(t, status="для тесту після 8–12 тижнів") for t in plan["meta"].get("tests", [])]
        names = self.T.get("symbol_names", {})
        ds, sy = self.T["description_symbols"], self.S["description_symbols"]
        known = sorted((n for n in dict.fromkeys(names.get(s, s) for s in ds["allowed"]) if n in sy), key=lambda n: -sy[n])
        parts = [f"{n} у {sy[n]}" + (f" з {self.N}" if i == 0 else "") for i, n in enumerate(known)]
        leader = tz.host(self.comps[0]["url"]).split(".")[0].replace("-reisen", "")
        self.sym_note = (f"Символи: у description не більше {['', 'одного', 'двох', 'трьох'][ds['max_distinct']]} різних простих символів ({' '.join(ds['allowed'])}), "
                         f"кожен не більше {['', 'одного разу', 'двох разів', 'трьох разів'][ds['max_each']]}, без кольорових емодзі; у title символів немає. "
                         f"У конкурентів у description: {', '.join(parts)}" + (f"; у лідера {leader} символів немає." if not self.comps[0]["description_symbols"] else "."))
        self.terms = self.term_rows()
        self.dev = self.developer()

    # ---------- data ----------
    def term_rows(self):
        tr, rows = self.plan.get("terms", {}), []
        for t in self.E["terms"]:
            p = tr.get(t["term"].lower())
            if not p:
                continue
            rows.append({"term": p.get("name", t["term"]), "uk": p["uk"], "df": t["competitors"], "sim": t["sim_best"], "where": p["where"],
                         "note": ("головне слово сторінки %s: лише анкор посилання" % t["other_page"]) if t["head_word_of_other_page"] else ""})
        return sorted(rows, key=lambda r: (-r["df"], -r["sim"]))

    def link_rows(self):
        st = {n["page"]: n["state"] for n in self.ctx["neighbours"]}
        rows = [[l["page"], st.get(l["page"], "?"), "; ".join(l["anchors"]), l["where"]] for l in self.plan.get("links", [])]
        return rows + [[l["page"], l.get("state", "існує"), "; ".join(l["anchors"]), l["where"]] for l in self.plan.get("links_extra", [])]

    def developer(self):
        """everything the text cannot solve: collected from the page analysis and stage-1 notes, plus the plan's own items"""
        own, a, out = self.own, self.plan["meta"]["A"], []
        if own.get("ok"):
            if own["title"] != a["title"]["text"]:
                out.append(f"Title: замінити на «{a['title']['text']}» ({len(a['title']['text'])} симв.). Зараз: «{own['title']}» ({own['title_len']} симв.).")
            if own["description"] != a["description"]["text"]:
                out.append(f"Description: замінити на «{a['description']['text']}» ({len(a['description']['text'])} симв.). Зараз {own['description_len']} симв."
                           + (" і символи " + " ".join(own["description_symbols"]) + ", яких правила не дозволяють." if set(own["description_symbols"]) - set(self.T["description_symbols"]["allowed"]) else "."))
            if " | ".join(own["h1"]) != a["h1"]["text"]:
                out.append(f"H1: замінити на «{a['h1']['text']}». Зараз: «{' | '.join(own['h1']) or 'немає'}».")
        else:
            out.append(f"Сторінка нова: задати title «{a['title']['text']}», description і H1 з комплекту A (розділ 10).")
        out += self.plan.get("developer", [])
        notes = self.ctx["decisions"].get("page_notes", {})
        for p, note in notes.items():
            if re.search("КАНІБАЛІЗАЦ", note, re.I) and (p == self.ctx["path"] or self.ctx["path"] in note):
                out.append(f"Канібалізація з етапу 1 ({p}): " + re.sub(r"^⚠\s*КАНІБАЛІЗАЦІЯ:\s*", "", note))
        if own.get("ok"):
            miss = [t for t in ("FAQPage", "BreadcrumbList") if t not in own["schema_types"]]
            # the plan may already carry its own JSON-LD item: then the automatic one would repeat it
            if miss and not any(re.search(r"json-?ld", d, re.I) for d in self.plan.get("developer", [])):
                out.append("JSON-LD: додати розмітку " + ", ".join(miss) + (" (зараз на сторінці розмітки JSON-LD не знайдено)." if not own["schema_types"] else f" (зараз є: {', '.join(own['schema_types'])})."))
            if own.get("ui_messages"):
                out.append(f"У HTML сторінки є {len(own['ui_messages'])} службових повідомлень, які бачить пошуковий робот, наприклад: «" + own["ui_messages"][0][:110]
                           + "…». Перевірити, чи вони приховані від індексації і чи працює лістинг.")
            dup = own["images"].get("duplicate_alt")
            if dup:
                out.append(f"Фото: {dup['count']} зображень мають однаковий alt «{dup['alt']}» — задати унікальні alt із ТЗ (розділ 4).")
            have = {tz.page_path(l["href"]) for l in own.get("internal_links", [])}
            exist = [l["page"] for l in self.plan.get("links", []) if next((n for n in self.ctx["neighbours"] if n["page"] == l["page"]), {}).get("state", "").startswith("існує") and l["page"] not in have]
            if exist:
                out.append("Внутрішні посилання: сторінки вже існують, але хаб на них не посилається: " + ", ".join(exist) + ".")
        return out

    # ---------- xlsx ----------
    def xlsx(self, path):
        P, E, N, own = self.plan, self.E, self.N, self.own
        wb = Workbook()
        wb.remove(wb.active)
        kw = sorted(E["keywords"], key=lambda r: (-r["volume"], -r["sim_main"]))
        pk = P["keywords"]
        rows = [[k["keyword"], k["translation"], k["volume"], k["sim_main"], pk[k["keyword"]]["place"], k["tier"],
                 f"{k['closest_adjacent']} ({k['sim_closest_adjacent']:.2f})" if k["closest_adjacent"] else "—", pk[k["keyword"]]["count"], pk[k["keyword"]]["block"], pk[k["keyword"]]["why"]] for k in kw]
        th = E["thresholds"]
        sheet(wb, "Семантика", ["ключ", "переклад", "частотність", "близькість до головного", "рекомендоване місце (title / H1 / H2 / текст / FAQ)", "рівень за близькістю",
                               "найближчий суміжний ключ (близькість)", "к-ть входжень", "блок сторінки", "пояснення"], rows, [44, 44, 12, 12, 24, 18, 40, 10, 52, 62],
              note=f"Близькість — косинусна, модель {E['model']} (відкрита модель, не Google): орієнтир поруч із частотністю й аналізом ТОП, а не замість них. "
                   f"Пороги: ≥{th['core']:.2f} — ядро, {th['mid']:.2f}–{th['core']:.2f} — H2/текст, <{th['mid']:.2f} — FAQ/довгий хвіст. Де рекомендоване місце відрізняється від рівня за близькістю, причина — у колонці «пояснення».")
        rows = []
        for st in self.sets:
            for el, key, lim in (("Title", "title", f"до {self.T['title_max']}"), ("Description", "description", f"до {self.T['description_max']}"), ("H1", "h1", "—")):
                rows.append([st["name"], st["status"], el, st[key]["text"], len(st[key]["text"]), lim, st[key]["why"] + (" " + self.sym_note if key == "description" and st["status"] == "впроваджувати" else ""), st["condition"]])
        cn = P["meta"].get("current_notes", {})
        if own.get("ok"):
            rows += [["Зараз на сайті", "для порівняння", "Title", own["title"], own["title_len"], f"до {self.T['title_max']}", cn.get("title", ""), ""],
                     ["Зараз на сайті", "для порівняння", "Description", own["description"], own["description_len"], f"до {self.T['description_max']}", cn.get("description", ""), ""],
                     ["Зараз на сайті", "для порівняння", "H1", " | ".join(own["h1"]), len(" | ".join(own["h1"])), "—", cn.get("h1", ""), ""]]
        ws, r0 = sheet(wb, "Мета-теги", ["комплект", "статус", "елемент", "текст (німецькою)", "довжина", "ліміт", "пояснення", "коли застосовувати / умова тесту"], rows, [16, 22, 13, 78, 10, 9, 90, 70],
                       note="Перший комплект (A) — узгоджені title + description + H1, його впроваджувати. Решта комплектів — для тесту після 8–12 тижнів, кожен зі своєю умовою. "
                            + P["meta"].get("note", "") + " " + self.sym_note)
        for i, row in enumerate(rows, r0 + 1):
            if row[1] == "впроваджувати" or row[0].startswith("Зараз"):
                fill_row(ws, i, 8, OK_FILL if row[1] == "впроваджувати" else GREY)

        def pos(c):
            p = c.get("serp_positions") or {}
            return "; ".join(f"{k}: {v}" for k, v in sorted(p.items(), key=lambda x: (x[0] != self.main, x[1]))) or "—"

        def comp_row(c, label=None):
            return [c["url"], label or pos(c), c["title"], c["description"], " | ".join(c["h1"]), c["text_chars_nospace"],
                    f"{c['lists']['content']} (+{c['lists']['link_lists']} списків посилань)", c["tables"]["count"], f"{c['images']['content']} (alt: {c['images']['with_alt']})",
                    ("так, %d питань%s" % (len(c["faq"]["questions"]), ", FAQPage" if c["faq"]["schema_faqpage"] else "")) if c["faq"]["present"] else "ні",
                    ("так: " + ", ".join((c["regions"]["in_headings"] or c["regions"]["in_link_anchors"])[:8])) if c["regions"]["block"] else "ні",
                    ("так: " + "; ".join(c["seasons"]["headings"][:2])) if c["seasons"]["block"] else "ні",
                    c["h2_count"], c["h3_count"], c["videos"], c["prices"]["mentions"], c["reviews"]["mentions"], "так" if c["filters"]["search_form"] else "ні",
                    c["cta"]["count"], ", ".join(c["keywords_exact_any"]) or "—", ", ".join(c["schema_types"]) or "—",
                    "; ".join(x for x in [c.get("note"), ("замінює " + c["replaces"]) if c.get("replaces") else None] if x) or ""]
        rows = [comp_row(c) for c in self.comps]
        if own.get("ok"):
            r = comp_row(own, "власна сторінка")
            ui = sum(tz.nosp(x) for x in own.get("ui_messages", []))
            r[5] = f"{own['text_chars_nospace']}" + (f" (ще {ui} симв. — службові повідомлення пошуку, у текст не враховано)" if ui else "")
            r[-1] = "сторінка сайту для порівняння; відкрита з браузерним User-Agent"
            rows.append(r)
        for r in self.S["replaced"]:
            tried = "; ".join(f"{t['url']} — {'відкрилася' if t['ok'] else 'не відкрилася (' + str(t['reason']) + ')'}" for t in r["tried"]) or "запасних не лишилось"
            rows.append([r["failed"], "не відкрилася"] + [""] * 19 + [f"{r['reason']}. Заміна: {tried}"])
        S = self.S
        rows.append([f"МЕДІАНА ({N} конкурентів)", "", "", "", "", S["text_median"], int(S["lists_median"]), int(S["tables_median"]), S["images_median"], f"FAQ у {S['with_faq']} з {N}",
                     f"у {S['with_region_block']} з {N}", f"у {S['with_season_block']} з {N}", S["h2_median"], S["h3_median"], f"у {S['with_video']} з {N}", "", "", "", "", "", "", ""])
        ws, r0 = sheet(wb, "Конкуренти", ["URL", "позиція (ключ: місце)", "title", "description", "H1", "обсяг тексту (симв. без пробілів)", "списки", "таблиці", "фото", "FAQ",
                                         "блок регіонів", "блок сезонів", "H2", "H3", "відео", "згадок цін", "згадок відгуків", "форма пошуку", "CTA", "ключі сторінки (точне входження)", "JSON-LD", "примітка"],
                       rows, [52, 30, 44, 60, 36, 16, 16, 9, 14, 20, 40, 40, 6, 6, 7, 9, 10, 9, 7, 40, 30, 60])
        for i, row in enumerate(rows, r0 + 1):
            fill = OK_FILL if row[1].startswith("власна") else (GAP_FILL if row[1] == "не відкрилася" else (GREY if str(row[0]).startswith("МЕДІАНА") else None))
            if fill:
                fill_row(ws, i, 22, fill)
        rows = [[t["term"], t["uk"], f"{t['df']} з {N}", t["sim"], t["where"], t["note"]] for t in self.terms]
        sheet(wb, "Семантичні слова", ["термін", "переклад укр.", "у скількох конкурентів", "близькість (для сортування)", "де використати (H2 / текст / FAQ)", "примітка"], rows, [30, 38, 14, 16, 46, 50],
              note=f"Критерій відбору — термін є щонайменше у {th['term_min_competitors']} з {N} конкурентів. Близькість до головного й суміжних ключів — лише колонка для сортування, не фільтр: "
                   "відкрита модель дає високі значення тільки фразам зі словами головного ключа, а окремим іменникам і назвам курортів — низькі. "
                   "Навігацію, бренди конкурентів, cookie-тексти й загальні слова відсічено стоп-списком (references/stoplist скіла tz-travel).")
        rows, uk = [], P.get("subtopics_uk", {})
        for tid, t in sorted(self.topics.items(), key=lambda x: (x[0].startswith("block:"), -x[1]["competitors"])):
            if t["competitors"] < 2:
                continue
            r = next((r for r in E["subtopics"] if "subtopic:" + r["subtopic"] == tid), None)
            rows.append([t["name"], uk.get(t["name"], ""), " | ".join(r["headings"][:8]) if r else "—", f"{t['competitors']} з {N}", t["own"], "так" if t["gap"] else "ні", t["source"],
                         (f"слова теми: {', '.join(r['topic_words']) or '—'}; у тексті сторінки знайдено: {', '.join(r['topic_words_in_own_text']) or '—'}; "
                          f"найближчий власний заголовок: «{r['own_best_heading']}» ({r['own_best_sim']:.2f})") if r else ""])
        ws, r0 = sheet(wb, "Підтеми", ["підтема", "переклад укр.", "приклади заголовків конкурентів", "у скількох конкурентів", f"є на {self.site_host}", "прогалина (так/ні)", "джерело", "примітка"],
                       rows, [48, 36, 110, 12, 30, 12, 26, 70],
                       note=f"Підтеми — кластери H2/H3 конкурентів за близькістю ≥{th['topic']:.2f} ({E['headings_total']} заголовків {N} конкурентів, показано кластери з ≥2 конкурентами); "
                            "перед кластеризацією із заголовків прибрано слова головного ключа, щоб групувати за підтемою. Наприкінці — блоки сторінки, знайдені розбором HTML (FAQ, регіони, сезон). "
                            f"Прогалина = тема є щонайменше у {th['gap_min_competitors']} конкурентів і її немає на власній сторінці. «Є на сторінці» лише тоді, коли збігаються і заголовок (близькість ≥{th['topic']:.2f}), і слова теми в тексті.")
        for i, row in enumerate(rows, r0 + 1):
            if row[5] == "так":
                fill_row(ws, i, 8, GAP_FILL)
        save(wb.save, path)

    # ---------- docx ----------
    def docx(self, path):
        P, E, N, S, T, own = self.plan, self.E, self.N, self.S, self.T, self.own
        doc = Document()
        doc.styles["Normal"].font.name, doc.styles["Normal"].font.size = "Calibri", Pt(10.5)
        for s in doc.sections:
            s.left_margin = s.right_margin = Cm(1.8)
            s.top_margin = s.bottom_margin = Cm(1.6)

        def para(text, bold=False):
            p = doc.add_paragraph(text)
            p.runs[0].bold = bold
        bul = lambda text: doc.add_paragraph(text, style="List Bullet")

        def shade(cell, color):
            shd = OxmlElement("w:shd")
            shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), color)
            cell._tc.get_or_add_tcPr().append(shd)

        def table(header, rows, widths, gap_col=None):
            t = doc.add_table(rows=1, cols=len(header))
            t.style = "Table Grid"
            for j, h in enumerate(header):
                run = t.rows[0].cells[j].paragraphs[0].add_run(h)
                run.bold, run.font.size, run.font.color.rgb = True, Pt(9), RGBColor(0xFF, 0xFF, 0xFF)
                shade(t.rows[0].cells[j], "1F3864")
            for row in rows:
                cells = t.add_row().cells
                for j, v in enumerate(row):
                    cells[j].paragraphs[0].add_run(str(v)).font.size = Pt(9)
                    if gap_col is not None and GAP_TAG in str(row[gap_col]):
                        shade(cells[j], "FCE4D6")
                    elif gap_col is not None and REQ_TAG in str(row[gap_col]):
                        shade(cells[j], "DDEBF7")
            for row in t.rows:
                for j, w in enumerate(widths):
                    row.cells[j].width = Cm(w)
            doc.add_paragraph()

        doc.add_heading(P["doc_title"], 0)
        for h in P["header"]:
            para(h)
        doc.add_heading("1. Мета сторінки, аудиторія, інтент", 1)
        for b in P["goal"]:
            bul(b)

        doc.add_heading("2. Обсяг тексту", 1)
        vr = P["volume"]["range"]
        para(f"Рекомендований обсяг: {num(vr[0])}–{num(vr[1])} символів без пробілів (приблизно {P['volume']['words']} слів).", True)
        bul(f"Сума блоків зі структури в розділі 3 — {num(self.lo)}–{num(self.hi)} символів.")
        bul(f"Медіана {N} конкурентів — {num(S['text_median'])} символів без пробілів; розкид від {num(S['text_min'])} до {num(S['text_max'])}.")
        bul(f"Текстові лідери з ТОП-5 за обсягом мають {num(S['text_top5'][-1])}–{num(S['text_top5'][0])} символів.")
        for b in P["volume"]["why"]:
            bul(b)

        doc.add_heading("3. Структура H1–H3", 1)
        para(f"У колонці «у конкурентів» два види позначок. «{GAP_TAG} (N з M)» — тема є щонайменше у {self.cfg['gap_min_competitors']} з {N} конкурентів і її немає на {self.site_host}. "
             f"«{REQ_TAG} (кластер, частотність головного ключа)» — кластер ключів, який етап 1 призначив цій сторінці: він має власний H2 або H3 незалежно від кількості конкурентів. Обсяг H2 включає його H3. "
             "Порядок на сторінці: H1 → 1–2 речення → лістинг → увесь текст, починаючи з H2.")
        rows = []
        for s in P["structure"]:
            marks = []
            if s.get("topic") and s["level"] == "H2":
                t = sorted((self.topics[x] for x in pc.topic_ids(s)), key=lambda t: (not t["gap"], -t["competitors"]))[0]
                marks.append(f"{GAP_TAG} ({t['competitors']} з {N})" if t["gap"] else f"у конкурентів: {t['competitors']} з {N}")
            for cid in s.get("clusters", []):
                marks.append(f"{REQ_TAG} ({cid}, {num(self.ctx['clusters'][cid]['volume'])})")
            rows.append([LEVEL.get(s["level"], s["level"]), s["heading"], chars(s.get("chars")), "; ".join(marks), s["note"]])
        table(["рівень", "заголовок (німецькою)", "обсяг, симв. без пробілів", "у конкурентів", "що писати"], rows, [1.4, 4.8, 2.0, 3.8, 5.8], gap_col=3)
        gaps = [t for t in self.topics.values() if t["gap"]]
        para("Кластери етапу 1 і їх заголовки: " + "; ".join(f"{cid} «{c['head']}» ({num(c['volume'])}) → «" + "», «".join(s["heading"] for s in P["structure"] if cid in s.get("clusters", [])) + "»"
                                                            for cid, c in self.ctx["clusters"].items()) + ".")
        para("Прогалини за конкурентами: " + ("; ".join(f"«{t['name']}» ({t['competitors']} з {N}, {t['source']})" for t in gaps) or "немає")
             + f". Підтеми — кластери H2/H3 конкурентів за близькістю ≥{E['thresholds']['topic']:.2f} і блоки сторінки з розбору HTML (аркуш «Підтеми» у meta-{self.ctx['slug']}.xlsx).")
        for tid, why in P.get("gaps_declined", {}).items():
            para(f"Прогалину «{self.topics[tid]['name']}» у структуру не включено: {why}")

        doc.add_heading("4. Списки, таблиці, фото, FAQ", 1)
        doc.add_heading("Списки", 2)
        for b in P["lists"]:
            bul(b)
        bul(f"Змістовні списки є у {S['with_content_lists']} з {N} конкурентів; більше трьох списків на сторінку не потрібно.")
        doc.add_heading("Таблиці", 2)
        for b in P["tables"]:
            bul(b)
        doc.add_heading("Фото", 2)
        for b in P["photos"]:
            bul(b)
        bul(f"Медіана у конкурентів — {S['images_median']} зображень, але більшість із них — картки готелів у лістингу.")
        doc.add_heading("FAQ", 2)
        para(f"FAQ є у {S['with_faq']} з {N} конкурентів, розмітку FAQPage мають {S['with_faqpage']}. {['', 'Одне питання', 'Два питання', 'Три питання', 'Чотири питання', 'П`ять питань', 'Шість питань', 'Сім питань', 'Вісім питань'][len(P['faq'])].replace('`', chr(39)) if len(P['faq']) < 9 else str(len(P['faq'])) + ' питань'}"
             f", {P.get('faq_answer', 'відповідь 180–220 символів')}:")
        table(["№", "питання (німецькою)", "звідки питання / що врахувати"], [[i, q["q"], q["note"]] for i, q in enumerate(P["faq"], 1)], [0.8, 8.5, 8.5])

        doc.add_heading("5. Обов'язкові ключі", 1)
        para(P.get("keywords_intro", "Кількість — точні входження або природна форма з тими самими словами. Порядок — за частотністю."))
        kw, pk = sorted(E["keywords"], key=lambda r: (-r["volume"], -r["sim_main"])), P["keywords"]
        table(["ключ", "частотність", "близькість", "місце", "входжень", "блок"],
              [[k["keyword"], k["volume"], f"{k['sim_main']:.2f}", pk[k["keyword"]]["place"], pk[k["keyword"]]["count"], pk[k["keyword"]]["block"]] for k in kw], [4.6, 1.6, 1.5, 2.6, 1.5, 6.0])
        mk, ex = self.main.title(), S["main_keyword_exact_in_text"]
        mk = next((s["heading"][:len(self.main)] for s in P["structure"] if s["level"] == "H1"), mk)
        total = len(self.heads) + len(self.faqx) + len(self.textx)
        para(f"Точна форма «{mk}»: не більше {T['exact_limit']} входжень на всій сторінці (заголовки + текст + FAQ; title і description не рахуються). "
             f"У конкурентів у тексті вона трапляється {ex[0]}–{ex[-1]} разів, медіана {int(statistics.median(ex))}. Розподіл:", True)
        table(["де", "кількість", "місця"],
              [["заголовки", len(self.heads), "; ".join("«%s»" % h for h in self.heads)], ["FAQ", len(self.faqx), "; ".join("«%s»" % q for q in self.faqx) + " (у відповідях — 0)"],
               ["текст", len(self.textx), "; ".join(self.textx)], ["разом", total, f"ліміт {T['exact_limit']}"]], [2.5, 1.8, 13.5])
        if P.get("keywords_note"):
            para(P["keywords_note"])

        doc.add_heading("6. Семантичне ядро тексту", 1)
        para(f"Слова й словосполучення з текстів конкурентів. Колонка «у конкурентів» показує, на скількох із {N} сторінок термін трапляється. Кожен термін достатньо вжити 1–2 рази у вказаному місці; назви регіонів з H3 — 2–3 рази.")
        para(f"Примітка: близькість рахує відкрита модель {E['model']}, а не Google. Це орієнтир поруч із частотністю й аналізом ТОП, а не заміна їм. "
             f"Терміни відібрано за частотою: кожен є щонайменше у {E['thresholds']['term_min_competitors']} конкурентів; близькість наведено лише для сортування.")
        table(["термін", "переклад", "у конкурентів", "близькість", "де використати"], [[t["term"], t["uk"], f"{t['df']} з {N}", f"{t['sim']:.2f}", t["where"]] for t in self.terms], [3.8, 4.4, 1.9, 1.7, 6.0])

        doc.add_heading("7. Внутрішні посилання", 1)
        para("Кожна сторінка отримує 1–2 посилання з тексту; анкори чергувати, не повторювати той самий анкор двічі. Посилання на сторінки зі станом «нова» ставити після їх запуску.")
        table(["сторінка", "стан", "анкори (німецькою)", "де в тексті"], self.link_rows(), [5.0, 1.6, 6.0, 5.2])
        if P.get("links_note"):
            para(P["links_note"])

        doc.add_heading("8. Заборони", 1)
        if self.ctx["banned"]:
            bul("Не використовувати " + ", ".join("«%s»" % b["label"] for b in self.ctx["banned"]) + " (і споріднені форми) у title, H1, H2, H3 і вступі: це головні слова сусідніх сторінок. "
                "Ці слова — лише в анкорах посилань на відповідні сторінки й у реченні навколо анкора, не більше 2 згадок кожного на весь текст.")
        nk = [n["main_keyword"] for n in self.ctx["neighbours"] if n["main_keyword"]]
        if nk:
            bul("Не оптимізувати текст під головні ключі інших сторінок: " + ", ".join("«%s»" % k for k in nk) + ".")
        bul(f"Без переспаму: «{mk}» у точній формі — не більше {T['exact_limit']} разів на всій сторінці: {len(self.heads)} у заголовках (H1 і один H2), {len(self.faqx)} у питанні FAQ, "
            f"{len(self.textx)} у тексті (розподіл — у розділі 5). В інших заголовках, у відповідях FAQ і в alt фото точної форми немає; чергувати з " + (", ".join("«%s»" % x for x in P["diluted_forms"]) if P.get("diluted_forms") else "розбавленими формами") + ".")
        bul(f"Не змінювати формулювання заголовків зі структури так, щоб у них з'явилася точна форма «{mk}».")
        for b in P.get("bans", []):
            bul(b)
        bul("Не вигадувати ціни, рейтинги й відгуки. Ціни «ab … €» підставляє замовник із лістингу.")

        doc.add_heading("9. Що конкуренти роблять краще і як їх обійти", 1)
        para("Що вони роблять краще:", True)
        for b in P["competitors_better"]:
            bul(b)
        para("Як обійти:", True)
        for b in P["how_to_beat"]:
            bul(b)

        doc.add_heading("10. Мета-теги", 1)

        def meta_table(st):
            table(["елемент", "текст (німецькою)", "довжина", "пояснення"],
                  [["Title", st["title"]["text"], f"{len(st['title']['text'])} (до {T['title_max']})", st["title"]["why"]],
                   ["Description", st["description"]["text"], f"{len(st['description']['text'])} (до {T['description_max']})", st["description"]["why"] + (" " + self.sym_note if st["status"] == "впроваджувати" else "")],
                   ["H1", st["h1"]["text"], len(st["h1"]["text"]), st["h1"]["why"]]], [2.2, 6.4, 1.8, 7.4])
        doc.add_heading("Впроваджувати: комплект A", 2)
        para("Один узгоджений комплект: title, description і H1 починаються однаково. " + P["meta"].get("note", ""))
        meta_table(self.sets[0])
        if len(self.sets) > 1:
            doc.add_heading("Для тесту після 8–12 тижнів", 2)
            para("Не впроваджувати одразу. Кожен комплект пробувати лише за своєї умови й міняти title разом із парним H1; порівнювати CTR і позиції в GSC за 4 тижні до і після.")
            for st in self.sets[1:]:
                para(f"{st['name']}. Умова: {st['condition']}", True)
                meta_table(st)
        if own.get("ok"):
            doc.add_heading("Зараз на сайті (для порівняння)", 2)
            table(["елемент", "текст", "довжина"], [["Title", own["title"], own["title_len"]], ["Description", own["description"], own["description_len"]], ["H1", " | ".join(own["h1"]), len(" | ".join(own["h1"]))]], [2.2, 13.0, 2.6])

        doc.add_heading("11. Для розробника", 1)
        para("Те, чого текст не вирішує. Цей розділ — не для копірайтера.")
        for b in self.dev:
            bul(b)
        save(doc.save, path)


def save(fn, path):
    try:
        fn(path)
    except PermissionError:
        alt = re.sub(r"(\.\w+)$", r".locked\1", path)
        fn(alt)
        print(f"УВАГА: {path} відкритий в іншій програмі; записано {alt}. Закрийте файл і повторіть запуск.")


def claude_md(ctx, C, out_files):
    """journal line in "## tz-travel" of the working folder's CLAUDE.md (never of a parent folder; created when missing)"""
    comps = ", ".join(tz.host(c["url"]) for c in C["competitors"] if c.get("ok"))
    line = f"- {ctx['path']} — {datetime.date.today().isoformat()}; файли: {', '.join(os.path.basename(f) for f in out_files)}; конкуренти ({C['summary']['analysed_ok']}): {comps}"
    return tz.sp.journal_add(ctx["workdir"], "tz-travel", f"{ctx['path']} — ", line)


def main():
    ap = tz.args_page(argparse.ArgumentParser())
    ap.add_argument("--suffix", default="", help="напр. .new — не перезаписувати затверджені файли")
    ap.add_argument("--no-claude-md", action="store_true")
    a = ap.parse_args()
    ctx = tz.load_ctx(a.workdir, a.page)
    C, E = pc.load_data(ctx)
    if not os.path.exists(ctx["f"]["plan"]):
        tz.die(f"немає {ctx['f']['plan']}. Зробіть чернетку: plan_check.py --draft, заповніть її за references/rules.md.")
    plan = tz.rjson(ctx["f"]["plan"])
    err, warn = pc.check(ctx, plan, C, E)
    for w in warn:
        print("УВАГА:", w)
    if err:
        print("План не пройшов перевірку, файли не зібрано:")
        for e in err:
            print("  ПОМИЛКА:", e)
        sys.exit(1)
    b = Build(ctx, plan, C, E)
    fx = os.path.join(a.workdir, f"meta-{ctx['slug']}{a.suffix}.xlsx")
    fd = os.path.join(a.workdir, f"tz-copywriter-{ctx['slug']}{a.suffix}.docx")
    b.xlsx(fx)
    b.docx(fd)
    if not a.no_claude_md:
        claude_md(ctx, C, [fx, fd])
    print("записано:", fx, "|", fd)
    print(f"точні входження «{ctx['main']}»: заголовки {len(b.heads)}, FAQ {len(b.faqx)}, текст {len(b.textx)}, разом {len(b.heads) + len(b.faqx) + len(b.textx)} (ліміт {ctx['T']['exact_limit']})")
    print(f"обсяг: блоки {b.lo}–{b.hi}, діапазон {plan['volume']['range'][0]}–{plan['volume']['range'][1]}; медіана {C['summary']['text_median']}, лідери ТОП-5 {C['summary']['text_top5']}")
    a_ = plan["meta"]["A"]
    print(f"комплект A: title {len(a_['title']['text'])}, description {len(a_['description']['text'])}, H1 {len(a_['h1']['text'])}")
    print("кластери:", "; ".join(f"{cid} → " + " / ".join(s["heading"] for s in plan["structure"] if cid in s.get("clusters", [])) for cid in ctx["clusters"]))
    print("прогалини:", "; ".join(f"{t['name']} ({t['competitors']} з {t['of']})" for t in b.topics.values() if t["gap"]) or "немає")
    print("для розробника:", len(b.dev), "пунктів")


if __name__ == "__main__":
    main()
