#!/usr/bin/env python3
"""Stage 0 (free): parse the input, normalise domains, look up earlier checks and placements, check Labs languages,
create the run folder with input.json, stage0.json and an empty site-notes.json for the homepage review.

Input file (references/input-template.md):
  Акцептор: https://example.kz/
  Ринок: Казахстан
  Тематика: туризм            (optional, default: туризм)
  Рекомендувати: 6
  Бюджет: 2,3 $
  Домени:
  site.kz — 3000
  other.com
"""
import datetime, glob, json, os, re, sys
import bl_common as bl

IN_WORK = {"виконується", "відправлено запит", "новий"}


def parse_input(text):
    inp = {"domains": [], "notes": []}
    seen = {}
    in_domains = False
    for raw in text.splitlines():
        line = raw.strip()
        if not line or line.startswith("#"):
            continue
        m = re.match(r"^(Акцептор|Ринок|Тематика|Рекомендувати|Бюджет|Домени)\s*:\s*(.*)$", line, re.I)
        if m:
            k, v = m.group(1).lower(), m.group(2).strip()
            in_domains = k == "домени"
            if k == "акцептор":
                inp["acceptor"] = bl.norm_domain(v)
            elif k == "ринок":
                inp["market_name"] = v
            elif k == "тематика":
                inp["theme"] = v.lower()
            elif k == "рекомендувати":
                inp["recommend"] = int(re.sub(r"\D", "", v) or 0)
            elif k == "бюджет":
                inp["budget"] = float(re.sub(r"[^\d.,]", "", v).replace(",", ".") or 0)
            continue
        if not in_domains:
            continue
        m = re.match(r"^[-*•]?\s*(\S+)\s*(?:[—–-]+\s*)?(\d[\d\s.,]*)?", line)
        if not m or "." not in m.group(1):
            continue
        d = bl.norm_domain(m.group(1))
        p = float(re.sub(r"[\s]", "", m.group(2)).replace(",", ".")) if m.group(2) else None
        p = int(p) if p is not None and float(p).is_integer() else p
        if d in seen:
            old = seen[d]["price"]
            new = min([x for x in (old, p) if x is not None], default=None)
            inp["notes"].append(f"дублікат {d}: ціни {old} і {p}, взято {new}")
            seen[d]["price"] = new
        else:
            seen[d] = {"domain": d, "price": p}
            inp["domains"].append(seen[d])
    for k in ("acceptor", "market_name", "recommend", "budget"):
        if not inp.get(k):
            sys.exit(f"У вхідних даних немає поля «{k}» (див. references/input-template.md)")
    if not inp["domains"]:
        sys.exit("У вхідних даних немає доменів")
    inp.setdefault("theme", "туризм")
    if inp["theme"] not in bl.CFG["themes"]:
        sys.exit(f"Немає набору слів для тематики «{inp['theme']}»: додай його в scripts\\markets.json → themes")
    return inp


def previous_checks(workdir, domains, skip):
    out = {}
    try:
        from openpyxl import load_workbook
    except ImportError:
        return out
    for f in sorted(glob.glob(os.path.join(workdir, "donors-*.xlsx"))):
        if os.path.basename(f) == skip or os.path.basename(f).startswith("~$"):
            continue
        m = re.search(r"(\d{4}-\d{2}-\d{2})", os.path.basename(f))
        try:
            ws = load_workbook(f, read_only=True)["Зведена"]
        except Exception:
            continue
        rows = list(ws.iter_rows(values_only=True))
        head = [str(c or "") for c in rows[0]]
        if "домен" not in head or "вердикт" not in head:
            continue
        di, vi = head.index("домен"), head.index("вердикт")
        for r in rows[1:]:
            if r[di] in domains:
                out[r[di]] = {"date": m.group(1) if m else "", "verdict": r[vi], "file": os.path.basename(f)}
    return out


def placements(workdir, domains):
    files = sorted(f for f in glob.glob(os.path.join(workdir, "placements", "*.xlsx")) if not os.path.basename(f).startswith("~$"))
    if not files:
        return None, {}
    from openpyxl import load_workbook
    f = files[-1]
    rows = list(load_workbook(f, read_only=True).active.iter_rows(values_only=True))
    head = [str(c or "").strip().lower() for c in rows[0]]

    def col(*names):
        return next((i for i, h in enumerate(head) if any(n in h for n in names)), None)
    ci = {"date": col("дата"), "url": col("url"), "status": col("статус"), "price": col("вартість"), "cur": col("валюта")}
    out = {}
    for r in rows[1:]:
        if ci["url"] is None or not r[ci["url"]]:
            continue
        host = bl.norm_domain(str(r[ci["url"]]))
        for d in domains:
            if host == d or host.endswith("." + d):
                status = str(r[ci["status"]] or "") if ci["status"] is not None else ""
                out.setdefault(d, []).append({"date": str(r[ci["date"]])[:10] if ci["date"] is not None else "", "status": status,
                                              "price": r[ci["price"]] if ci["price"] is not None else None, "currency": r[ci["cur"]] if ci["cur"] is not None else "",
                                              "url": str(r[ci["url"]]), "in_work": status.strip().lower() in IN_WORK})
    return os.path.basename(f), out


def main():
    ap = bl.parser(__doc__.split("\n")[0])
    ap.add_argument("--input", required=True, help="файл із вхідними даними")
    a = ap.parse_args()
    inp = parse_input(open(a.input, encoding="utf-8").read())
    code, market = bl.find_market(inp["market_name"])
    inp["market"] = code
    today = a.today or datetime.date.today().isoformat()
    inp["date"] = today
    slug = inp["acceptor"].replace(".", "-")
    a.run = a.run or f"{slug}-{today}"
    inp["output"] = f"donors-{slug}-{today}.xlsx"
    run = bl.Run(a)
    for d in inp["domains"]:
        d["type"] = "ринок" if bl.is_market_tld(d["domain"], market) else "загальний"
    run.save("input.json", inp)
    run = bl.Run(a)
    doms = [d["domain"] for d in inp["domains"]]

    # Labs languages for the location (free reference endpoint, cached)
    if not run.cached("_labs", bl.EP_LANGS) and not run.offline:
        r = run.api("dataforseo_labs/locations_and_languages", None, method="GET")
        if r:
            run.write("_labs", bl.EP_LANGS, r)
    avail, _ = bl.labs_languages(run)
    notes = list(inp["notes"])
    if avail is None:
        langs = market["languages"]
        notes.append("довідник мов Labs не отримано: мови ринку взято з markets.json без перевірки")
    else:
        langs = [l for l in market["languages"] if l in avail]
        for l in market["languages"]:
            if l not in avail:
                notes.append(f"мови «{l}» для location {market['location_code']} у Labs немає — запити лише мовами: {', '.join(langs)}")
    prev = previous_checks(run.workdir, doms, inp["output"])
    pfile, placed = placements(run.workdir, doms)
    if pfile is None:
        notes.append("файлу розміщень у placements\\*.xlsx немає — стовпець «розміщувались раніше» заповнюється лише з site-notes.json (поле placed)")
    s0 = {"languages": langs, "labs_languages": avail, "previous": prev, "placements_file": pfile, "placements": placed, "notes": notes}
    run.save("stage0.json", s0)
    if not os.path.exists(run.path("site-notes.json")):
        run.save("site-notes.json", {d: {"status": None, "topic": "", "travel_rubric": None, "travel_url": "", "ad_section_url": "", "ad_listing_pattern": "", "language": "",
                                         "region": "", "audience": "", "aggregator": None, "articles": [], "placed": "", "manual_check": []} for d in doms})

    print(f"Запуск: {run.name}  ({run.dir})")
    print(f"Акцептор: {inp['acceptor']} | ринок: {market['names'][0]} (location {market['location_code']}, {market['se_domain']}) | мови Labs: {', '.join(langs)} "
          f"| тематика: {inp['theme']} | рекомендувати: {inp['recommend']} | бюджет: {inp['budget']:.2f} $")
    for d in inp["domains"]:
        p = prev.get(d["domain"])
        pl = placed.get(d["domain"])
        print(f"  {d['domain']:<26} {d['type']:<10} ціна: {d['price'] if d['price'] is not None else '—':<8} "
              f"раніше: {(p['date'] + ' ' + str(p['verdict'])) if p else 'ні':<22} розміщення: {('; '.join(x['date'] + ' ' + x['status'] + (' — ВЖЕ В РОБОТІ' if x['in_work'] else '') for x in pl)) if pl else 'ні'}")
    for n in notes:
        print("  ! " + n)
    print("\nКошторис (dry-run):")
    P1 = bl.plan_stage1(run)
    P2, cands = bl.plan_stage2(run)
    t, tr = bl.print_plan(P1 + P2, run.budget, run.spent())
    print(f"  етап 2 пораховано для {len(cands)} кандидатів (N+1); до етапу 1 склад умовний: спершу домени з TLD ринку")
    print("\nДалі: WebFetch головної кожного домену → заповнити site-notes.json; потім stage1.py --yes" if tr <= run.budget else
          "\nСТОП: кошторис із запасом перевищує бюджет — зменш кількість доменів або кандидатів, або підніми бюджет.")


if __name__ == "__main__":
    main()
