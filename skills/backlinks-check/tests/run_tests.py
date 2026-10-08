#!/usr/bin/env python3
"""Offline tests of the backlinks-check skill.  No paid request is sent: every script runs with --offline.

  1. dry-run estimate for the poehalisnami.kz list with an empty cache against the real spend of that run;
  2. the workbook built only from tests/fixtures/poehalisnami-kz-2026-10-08 against the accepted numbers;
  3. a partial cache (poehalisnami.ua in <repo>/backlinks/cache): the estimate lists what is missing and its cost.

Run:  python -X utf8 tests/run_tests.py [--keep]
"""
import json, os, re, shutil, subprocess, sys, tempfile

TESTS = os.path.dirname(os.path.abspath(__file__))
SKILL = os.path.dirname(TESTS)
SCRIPTS = os.path.join(SKILL, "scripts")
FX_KZ = os.path.join(TESTS, "fixtures", "poehalisnami-kz-2026-10-08")
FX_UA = os.path.join(TESTS, "fixtures", "poehalisnami-ua-2026-10-08")
REAL_CACHE = os.path.join(os.path.dirname(os.path.dirname(SKILL)), "backlinks", "cache")
TODAY = "2026-10-08"
ACTUAL_SPEND_KZ = 1.77          # by balance, run of 2026-10-08 (1,765 $ + 0,003 $ for one SERP task re-sent later)
CANDIDATES_KZ = "ng.kz,7152.kz,iapn.kz,lifepvl.kz,tourlib.net,socportal.info,anydaylife.com"
CANDIDATES_UA = "rayon.in.ua,konkurent.ua,ukrlib.com.ua,tviysvit.com.ua,kg.ua,vinnitsa.info"


def run(script, *args, ok=(0,)):
    p = subprocess.run([sys.executable, "-X", "utf8", os.path.join(SCRIPTS, script), *args], capture_output=True, text=True, encoding="utf-8")
    if p.returncode not in ok:
        print(p.stdout[-3000:], p.stderr[-3000:])
        raise SystemExit(f"{script} завершився з кодом {p.returncode}")
    return p.stdout + p.stderr


def test1(tmp):
    """Estimate with an empty cache (only the free Labs reference list) vs the real spend."""
    work, cache = os.path.join(tmp, "t1"), os.path.join(tmp, "t1cache")
    os.makedirs(work)
    shutil.copytree(os.path.join(FX_KZ, "cache", "_labs"), os.path.join(cache, "_labs"))
    run("stage0.py", "--workdir", work, "--cache-dir", cache, "--input", os.path.join(FX_KZ, "run", "input.md"), "--run", "t", "--today", TODAY, "--offline")
    out = run("estimate.py", "--workdir", work, "--cache-dir", cache, "--run", "t", "--today", TODAY, "--offline")
    m = re.search(r"РАЗОМ до запиту: ([\d.]+) \$ без запасу, ([\d.]+) \$ із запасом", out)
    plain, reserve = float(m.group(1)), float(m.group(2))
    print("\nТЕСТ 1 — dry-run для списку poehalisnami.kz (порожній кеш)")
    print("\n".join(l for l in out.splitlines() if re.match(r"^(етап|[12] |  етап|  РАЗОМ)", l)))
    res = True
    for name, v in (("без запасу", plain), ("із запасом 15%", reserve)):
        dev = (v - ACTUAL_SPEND_KZ) / ACTUAL_SPEND_KZ * 100
        good = abs(dev) <= 20
        res &= good
        print(f"  кошторис {name}: {v:.4f} $ проти факту {ACTUAL_SPEND_KZ:.2f} $ → {dev:+.1f}%  {'OK' if good else 'ПОЗА ±20%'}")
    return res


def test2(tmp):
    """Stages 0-2 and the workbook only from fixtures; numbers against expected.json."""
    from openpyxl import load_workbook
    work, cache = os.path.join(tmp, "t2"), os.path.join(FX_KZ, "cache")
    os.makedirs(work)
    common = ["--workdir", work, "--cache-dir", cache, "--run", "kz", "--today", TODAY, "--offline"]
    run("stage0.py", *common, "--input", os.path.join(FX_KZ, "run", "input.md"))
    for f in ("site-notes.json", "verdicts.json"):
        shutil.copyfile(os.path.join(FX_KZ, "run", f), os.path.join(work, "runs", "kz", f))
    run("stage1.py", *common)
    s2 = run("stage2.py", *common, "--candidates", CANDIDATES_KZ)
    out_xlsx = os.path.join(work, "donors-test.xlsx")
    rep = run("build_report.py", *common, "--out", out_xlsx, "--no-journal")
    led = os.path.join(work, "runs", "kz", "ledger.json")
    paid = json.load(open(led, encoding="utf-8")) if os.path.exists(led) else []
    exp = json.load(open(os.path.join(FX_KZ, "expected.json"), encoding="utf-8"))
    wb = load_workbook(out_xlsx)
    rows = list(wb["Зведена"].iter_rows(values_only=True))
    head = list(rows[0])
    COL = {"recommendation": "рекомендація", "verdict": "вердикт", "price": "ціна, грн", "price_per_1k_market": "ціна за 1K трафіку ринку, грн", "price_per_1k_total": "ціна за 1K загального трафіку, грн",
           "rd": "домени-донори", "rank": "ранг", "spam score": "spam score", "spam": "spam score", "market_etv": "трафік/міс ринку", "total_etv": "загальний трафік", "market_share": "частка ринку, %",
           "top10": "ключів у ТОП-10 (ринок)"}
    got = {r[0]: r for r in rows[1:] if r[0]}
    diffs, checked = [], 0
    for d, e in exp.items():
        if d.startswith("_"):
            continue
        for k, want in e.items():
            have = got[d][head.index(COL[k])]
            if want is None:
                have = have if isinstance(have, (int, float)) else None
            checked += 1
            if k == "recommendation":
                want, have = want.split(" — ")[0], (have or "").split(" — ")[0]
            if have != want:
                diffs.append((d, COL[k], want, have))
    total_row = got.get("РАЗОМ")
    total = total_row[5] if total_row else None
    print("\nТЕСТ 2 — xlsx лише з fixtures, без платних запитів")
    print(f"  аркуші: {', '.join(wb.sheetnames)}")
    print(f"  платних запитів у ledger: {len(paid)}; порівняно значень: {checked}; розбіжностей: {len(diffs)}")
    for d, col, want, have in diffs:
        print(f"    {d} | {col}: у прийнятому файлі {want!r}, зі скіла {have!r}")
    print(f"  вартість рекомендованих: {total} грн (у прийнятому файлі {exp.get('_total_price')})")
    print("  розбіжності verdict / verdict_auto і рекомендацій (з виводу build_report):")
    blk = rep.split("матеріал для калібрування rules.md:")[1].split("Журнал:")[0] if "матеріал для калібрування" in rep else ""
    for l in blk.strip().splitlines():
        print("  " + l[:230])
    return not diffs and not paid and total == exp.get("_total_price")


def test3(tmp):
    """Partial cache: the estimate must list what is missing and price it, without asking the API."""
    print("\nТЕСТ 3 — частковий кеш poehalisnami.ua")
    if not os.path.isdir(os.path.join(REAL_CACHE, "rayon.in.ua")):
        print(f"  ПРОПУЩЕНО: немає робочого кешу {REAL_CACHE} (він не в git). Тест має сенс на ПК, де виконувався запуск 2026-10-08.")
        return None
    work = os.path.join(tmp, "t3")
    os.makedirs(work)
    common = ["--workdir", work, "--cache-dir", REAL_CACHE, "--run", "ua", "--today", TODAY, "--offline"]
    run("stage0.py", *common, "--input", os.path.join(FX_UA, "input.md"))
    out = run("estimate.py", *common, "--candidates", CANDIDATES_UA, "--missing")
    print("\n".join("  " + l for l in out.splitlines() if re.match(r"^(етап|[12] |  етап|  РАЗОМ|  бракує|  повна|  платних)", l)))
    miss = [l for l in out.splitlines() if re.match(r"^    \S+\s+\S+\s+[\d.]+ \$", l)]
    print(f"  перші з {len(miss)} відсутніх відповідей:")
    for l in miss[:8]:
        print("  " + l.rstrip())
    s1 = run("stage1.py", *common, ok=(0, 1))
    refused = "платних запитів не виконано" in s1
    led = os.path.join(work, "runs", "ua", "ledger.json")
    print(f"  stage1.py --offline з неповним кешем: {'відмовився запитувати і показав, чого бракує' if refused else 'НЕ зупинився'}; ledger: {'порожній' if not os.path.exists(led) else 'Є ЗАПИСИ'}")
    m = re.search(r"РАЗОМ до запиту: ([\d.]+) \$", out)
    return bool(miss) and refused and not os.path.exists(led) and float(m.group(1)) > 0


def main():
    keep = "--keep" in sys.argv
    tmp = tempfile.mkdtemp(prefix="backlinks-check-tests-")
    try:
        results = {"1": test1(tmp), "2": test2(tmp), "3": test3(tmp)}
    finally:
        if keep:
            print("\nтимчасова папка:", tmp)
        else:
            shutil.rmtree(tmp, ignore_errors=True)
    print("\nПідсумок: " + "; ".join(f"тест {k}: {'OK' if v else ('пропущено' if v is None else 'НЕ ПРОЙДЕНО')}" for k, v in results.items()))
    sys.exit(0 if all(v is not False for v in results.values()) else 1)


if __name__ == "__main__":
    main()
