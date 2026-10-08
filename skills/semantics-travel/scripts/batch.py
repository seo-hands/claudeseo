#!/usr/bin/env python3
"""semantics-travel / batch: several pages from a CSV (url;seed1|seed2|...) with a shared keyword pool and shared SERP cache.

python batch.py plan   pages.csv [--workdir .]            parse the CSV, show pages/seeds and the combined cost estimate (nothing requested)
python batch.py run    pages.csv [--workdir .] [--yes]    collect per page -> shared pool (dedupe between pages) -> SERP once per unique keyword -> analyze each page
python batch.py summary [--workdir .]                     recluster all pages (NO API calls) and rebuild batch-summary.xlsx

Layout (workdir = batch root, CLAUDE.md with the "## semantics-travel" section here):
  pages/<slug>/      one normal project per page (own keywords.json, decisions, outputs); its CLAUDE.md is generated from the root one
                     with page_base/slug taken from the URL
  serp-raw-regular/  SHARED SERP cache (a keyword is requested once for the whole batch)
  batch-summary.xlsx page | seeds | clusters | volume | cannibalization ; sheet "Канібалізація"; sheet "Пул ключів"
Pool rule: a keyword stays on ONE page (where the site already ranks for it, else the page whose seed words match best, else the first page);
the other pages get it in "Відфільтровані" (дублікат між сторінками) and it is listed as a cannibalization candidate.
Cannibalization is also reported when keywords of different pages share >= 4 TOP-10 URLs, and when two pages end with the same target page.
"""
import argparse, csv, collections, json, os, re, subprocess, sys
from urllib.parse import urlsplit

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import sp_common as C

HERE = os.path.dirname(os.path.abspath(__file__))


def read_csv(path):
    rows = []
    with open(path, encoding="utf-8-sig", newline="") as f:
        for r in csv.reader(f, delimiter=";"):
            if len(r) >= 2 and r[0].strip().lower() not in ("url", ""):
                rows.append({"url": r[0].strip(), "seeds": [s.strip() for s in r[1].split("|") if s.strip()]})
    return rows


def page_slug(url):
    p = urlsplit(url).path.strip("/")
    return re.sub(r"[^a-z0-9]+", "-", p.lower()).strip("-") or "home"


def make_page_dir(root, S, row):
    d = os.path.join(root, "pages", page_slug(row["url"]))
    os.makedirs(d, exist_ok=True)
    src = open(os.path.join(root, "CLAUDE.md"), encoding="utf-8").read()
    path = urlsplit(row["url"]).path.rstrip("/") or "/"
    src = re.sub(r"(?m)^(\s*-?\s*)page_base:.*$", rf"\1page_base: {path}", src)
    src = re.sub(r"(?m)^(\s*-?\s*)slug:.*$", rf"\1slug: {page_slug(row['url'])}", src)
    open(os.path.join(d, "CLAUDE.md"), "w", encoding="utf-8").write(src)
    return d


def run(cmd):
    print("$", " ".join(cmd[1:4]), "...")
    r = subprocess.run([sys.executable, "-X", "utf8"] + cmd, capture_output=True)
    sys.stdout.write((r.stdout or b"").decode("utf-8", "replace")[-1500:])
    if r.returncode:
        sys.stderr.write((r.stderr or b"").decode("utf-8", "replace")[-800:])
        sys.exit(r.returncode)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("cmd", choices=["plan", "run", "summary"])
    ap.add_argument("csv", nargs="?")
    ap.add_argument("--workdir", default=".")
    ap.add_argument("--competitors", type=int, default=3)
    ap.add_argument("--yes", action="store_true")
    a = ap.parse_args()
    root = os.path.abspath(a.workdir)
    try:
        S = C.load_settings(root)
    except C.ConfigMissing as e:
        sys.exit(C.config_help(e))
    rows = read_csv(a.csv) if a.csv else []
    cache = os.path.join(root, "serp-raw-regular")

    if a.cmd == "plan":
        if not rows:
            sys.exit("CSV порожній: потрібні колонки url;seed (кілька seed через |)")
        est_rows = []
        for r in rows:
            n = len(r["seeds"])
            print(f"- {r['url']}  seeds: {r['seeds']}")
            est_rows += [("ranked_keywords сторінки", 1, C.PRICES["labs_ranked_keywords"]), ("keyword_suggestions", n, C.PRICES["labs_keyword_suggestions"]),
                         ("related_keywords", n, C.PRICES["labs_related_keywords"]), ("SERP seed (конкуренти)", 1, C.PRICES["serp_live_regular"]),
                         ("ranked_keywords конкурентів", a.competitors, C.PRICES["labs_ranked_keywords"]), ("bulk KD", 1, C.PRICES["labs_bulk_keyword_difficulty"])]
        total = C.print_estimate([(f"{lab}", n, p) for lab, n, p in est_rows])
        up = len(rows) * S.max_keywords
        print(f"SERP для унікальних ключів: до {up} × ${C.PRICES['serp_live_regular']} = ${up * C.PRICES['serp_live_regular']:.2f} (спільний кеш: дублікати між сторінками й уже зібрані не запитуються)")
        print(f"Орієнтовно разом до ${total + up * C.PRICES['serp_live_regular']:.2f}. Підтвердіть, тоді: batch.py run {a.csv} --yes")
        return

    if a.cmd == "run":
        if not a.yes:
            sys.exit("Платний режим: спочатку `plan`, покажіть вартість користувачу, дочекайтесь підтвердження і запустіть з --yes.")
        os.makedirs(cache, exist_ok=True)
        dirs = []
        for r in rows:                                          # 1. collect per page
            d = make_page_dir(root, S, r)
            dirs.append((r, d))
            run([os.path.join(HERE, "collect.py"), "--workdir", d, "--url", r["url"], "--seeds"] + r["seeds"] + ["--competitors", str(a.competitors), "--yes"])
        pool(root, dirs)                                         # 2. shared pool
        for r, d in dirs:                                        # 3. SERP once per unique keyword (shared cache)
            run([os.path.join(HERE, "fetch_serp.py"), "--workdir", d, "--cache", cache, "--yes"])
        a.cmd = "summary"

    if a.cmd == "summary":
        dirs = sorted(os.path.join(root, "pages", x) for x in os.listdir(os.path.join(root, "pages"))) if os.path.isdir(os.path.join(root, "pages")) else []
        for d in dirs:
            run([os.path.join(HERE, "analyze.py"), "--workdir", d, "--serp-dir", cache])
        summary(root, dirs, cache)


def pool(root, dirs):
    """Keep every keyword on one page; duplicates go to `filtered` of the other pages."""
    kws = {}
    for r, d in dirs:
        kw = json.load(open(os.path.join(d, "keywords.json"), encoding="utf-8"))
        for k in kw["keywords"]:
            kws.setdefault(k["keyword"], []).append((d, r, k))
    owner = {}
    for kw, lst in kws.items():
        if len(lst) == 1:
            continue
        def score(item):
            d, r, k = item
            seedwords = {w for s in r["seeds"] for w in re.findall(r"\w{4,}", s.lower())}
            return (1 if k.get("own_position") else 0, sum(1 for w in seedwords if w in kw.lower()))
        owner[kw] = max(lst, key=score)[0]
    for r, d in dirs:
        fn = os.path.join(d, "keywords.json")
        kw = json.load(open(fn, encoding="utf-8"))
        keep, drop = [], []
        for k in kw["keywords"]:
            (keep if owner.get(k["keyword"], d) == d else drop).append(k)
        kw["keywords"] = keep
        kw["filtered"] += [{"keyword": k["keyword"], "reason": f"дублікат між сторінками (залишено на {os.path.basename(owner[k['keyword']])})"} for k in drop]
        json.dump(kw, open(fn, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    json.dump({k: [x[0] for x in v] for k, v in kws.items() if len(v) > 1}, open(os.path.join(root, "batch-pool-duplicates.json"), "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    print(f"пул: {len(kws)} унікальних ключів, дублікатів між сторінками: {len(owner)}")


def summary(root, dirs, cache):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    dup = json.load(open(os.path.join(root, "batch-pool-duplicates.json"), encoding="utf-8")) if os.path.exists(os.path.join(root, "batch-pool-duplicates.json")) else {}
    pages, targets, kwpage = [], collections.defaultdict(list), {}
    for d in dirs:
        cj = json.load(open(os.path.join(d, "clusters.json"), encoding="utf-8"))
        name = os.path.basename(d)
        vol = sum(c["total"] for c in cj["clusters"])
        for c in cj["clusters"]:
            targets[c["page"]].append(name)
            for m in c["members"]:
                kwpage[m] = name
        kj = json.load(open(os.path.join(d, "keywords.json"), encoding="utf-8"))
        pages.append([name, "; ".join(kj.get("meta", {}).get("seeds", [])), len(cj["clusters"]), vol, len(kj["keywords"]), 0])
    cann = []
    for kw, ds in dup.items():
        cann.append(["той самий ключ у кількох сторінках", kw, ", ".join(os.path.basename(x) for x in ds), "залишено на одній, інші — «дублікат між сторінками»"])
    for tp, names in targets.items():
        if len(set(names)) > 1:
            cann.append(["одна цільова сторінка для різних вхідних сторінок", tp, ", ".join(sorted(set(names))), "об'єднати або розвести інтент"])
    # SERP overlap between keywords of different pages
    import glob
    S = {}
    for f in glob.glob(os.path.join(cache, "*.json")):
        j = json.load(open(f, encoding="utf-8"))
        if j["keyword"] in kwpage:
            S[j["keyword"]] = {C.nu(it["url"]) for it in j["response"]["tasks"][0]["result"][0]["items"] if it["type"] == "organic"}
    ks = sorted(S)
    for i, x in enumerate(ks):
        for y in ks[i + 1:]:
            if kwpage[x] != kwpage[y] and len(S[x] & S[y]) >= 4:
                cann.append(["схожа видача (≥4 спільних URL)", f"{x} ↔ {y}", f"{kwpage[x]} / {kwpage[y]}", f"{len(S[x] & S[y])} спільних URL"])
    for p in pages:
        p[5] = sum(1 for c in cann if p[0] in c[2])
    wb = Workbook()
    ws = wb.active
    ws.title = "Сторінки"
    ws.append(["сторінка", "seeds", "кластерів", "обсяг", "ключів", "канібалізація (записів)"])
    for p in pages:
        ws.append(p)
    ws2 = wb.create_sheet("Канібалізація")
    ws2.append(["тип", "ключ / сторінка", "сторінки", "дія"])
    for c in cann:
        ws2.append(c)
    ws3 = wb.create_sheet("Пул ключів")
    ws3.append(["ключ", "сторінка-власник"])
    for k, p in sorted(kwpage.items()):
        ws3.append([k, p])
    for w in (ws, ws2, ws3):
        for c in w[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="1F4E78")
    target = os.path.join(root, "batch-summary.xlsx")
    try:
        wb.save(target)
    except PermissionError:
        target = os.path.join(root, "batch-summary.new.xlsx")
        wb.save(target)
        print("УВАГА: xlsx відкритий в Excel, збережено як", target)
    print(f"batch-summary: {len(pages)} сторінок, {len(cann)} записів канібалізації -> {target}")


if __name__ == "__main__":
    main()
