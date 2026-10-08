#!/usr/bin/env python3
"""semantics-travel / analyze: lemma clusters + page choice by competitor landing-page types (mode recluster; NO API calls).

python analyze.py [--workdir .] [--claude-md CLAUDE.md] [--out-dir .] [--serp-dir serp-raw-regular]

Inputs  (workdir): keywords.json, serp-raw-regular/*.json, semantics-decisions.json (optional), *-demand.json (region SERPs, optional),
                   landing-verification.json (optional), existing-pages.json (optional)
Outputs (out-dir): serp-data.json, semantics-<slug>.xlsx, cluster-map.html, clusters.json
Rules   : presets/travel.py (engine) + scripts/profiles/<market>.json (language/country words); project overrides: semantics-profile.json, semantics-rules.py;
          settings from CLAUDE.md; manual decisions from semantics-decisions.json.

Page choice
  every organic URL of a keyword's TOP is classified (hub / region hub / last minute / all inclusive / pauschalreise / hotel / info / rundreise ...)
  and mapped to a candidate page; types of one product are counted together (country + region pages).
    >= thr_ok (5)  -> recommended page of that type;  thr_disputed..4 -> "спірно" (verified in landing-verification.json / contested verdicts);
    <= 2 -> no separate page, the keyword goes to the broader (parent) page.
  Lemma clusters are the base of grouping and are split only where their keywords land on different pages.
"""
import argparse, collections, glob, importlib.util, json, os, re, sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import sp_common as C

ITR = {"commercial": "комерційний", "transactional": "транзакційний", "informational": "інформаційний", "navigational": "навігаційний", None: "н/д"}


def is_ext(k):
    s = k.get("source", "")
    return s.startswith("region:")


def ext_region(k):
    s = k["source"]
    return s.split(":", 1)[1] if ":" in s else s


def ext_group(k):
    return k.get("ext_group")


def load_data(workdir, serp_dir):
    kw = json.load(open(os.path.join(workdir, "keywords.json"), encoding="utf-8"))
    dec = C.load_decisions(workdir)
    C.apply_decisions_to_keywords(kw, dec)
    K = {k["keyword"]: k for k in kw["keywords"] if not is_ext(k)}
    serp = {}
    sd = serp_dir if os.path.isabs(serp_dir) else os.path.join(workdir, serp_dir)
    for f in glob.glob(os.path.join(sd, "*.json")):
        d = json.load(open(f, encoding="utf-8"))
        res = d["response"]["tasks"][0]["result"][0]
        rows = [{"position": it["rank_group"], "url": it["url"], "domain": it["domain"], "title": it.get("title") or "", "snippet": it.get("description") or ""}
                for it in res["items"] if it["type"] == "organic"]
        serp[d["keyword"]] = {"results": rows, "request": d["request"], "fetched_at": d["fetched_at"], "api_datetime": res.get("datetime"), "check_url": res.get("check_url")}
    missing = [k for k in K if k not in serp]
    if missing:
        sys.exit(f"Немає сирого SERP для {len(missing)} ключів (запустіть fetch_serp.py): {missing[:5]}")
    return kw, dec, K, serp


def serp_cluster(order, S, hard=None, soft=None):
    cl = []
    for k in order:
        best, bs = None, -1
        for c in cl:
            if hard:
                ov = [len(S[k] & S[m]) for m in c]
                ok, score = min(ov) >= hard, sum(ov) / len(ov)
            else:
                ov0 = len(S[k] & S[c[0]])
                ok, score = ov0 >= soft, ov0
            if ok and score > bs:
                best, bs = c, score
        if best is None:
            cl.append([k])
        else:
            best.append(k)
    return cl


def run(workdir, S, args):
    R, rules_path = C.load_rules(S, workdir)
    kwj, DEC, K, serp = load_data(workdir, args.serp_dir)
    out = args.out_dir or workdir
    os.makedirs(out, exist_ok=True)
    BASE, PB = S.site, S.page_base.rstrip("/")
    THR_OK, THR_DISPUTED = S.thr_ok, S.thr_disputed
    EXISTING = set(S.existing_pages)
    ef = os.path.join(workdir, "existing-pages.json")
    if os.path.exists(ef):
        EXISTING |= set(json.load(open(ef, encoding="utf-8")))
    KEYWORD_OVERRIDES = {k: (v["page"], v.get("reason", "")) for k, v in DEC["keyword_overrides"].items()}
    VERDICTS = {p: (v["verdict"], v.get("reason", "")) for p, v in DEC["contested_verdicts"].items()}
    MANUAL = DEC["manual_check"]
    NOTES = dict(R.GENERIC_NOTES)
    NOTES.update(DEC["page_notes"])
    RULES = list(S.business_rules) + [r for r in DEC["business_rules"]]
    LABS_POS = DEC["known_positions"]

    order = sorted(K, key=lambda k: -K[k]["volume"])
    Sset = {k: {C.nu(r["url"]) for r in v["results"]} for k, v in serp.items() if k in K}
    main_kw = S.get("main_keyword") or order[0]
    chk = DEC["serp_manual_check"]
    if chk and chk.get("keyword") in serp:
        hit = sum(1 for r in serp[chk["keyword"]]["results"] if any(x in r["domain"] for x in chk["domains"]))
        print(f"перевірка «{chk['keyword']}»: {hit} із {len(chk['domains'])} доменів збігаються з ручним ТОП ({args.serp_dir})")
    elif main_kw in serp:
        print(f"Звірте видачу вручну: {main_kw} -> {serp[main_kw]['check_url']}")

    # ---- classify every URL ----
    URLS, kw_types, kw_pages, kw_n = [], {}, {}, {}
    for k in order:
        tc, pc = collections.Counter(), collections.Counter()
        kwreg, topic = R.detect_region(k.lower()), R.topic_for(k)
        for r in serp[k]["results"]:
            typ, region = R.classify_url(r["url"], r["title"])
            page = R.page_for(typ, region, kwreg, topic)
            r["type"], r["region"], r["page"] = typ, region, page
            tc[R.type_label(typ, region)] += 1
            if page:
                pc[page] += 1
            URLS.append([k, r["position"], r["url"], r["domain"].replace("www.", ""), R.type_label(typ, region), region or "", page or "—"])
        kw_types[k], kw_pages[k], kw_n[k] = tc, pc, len(serp[k]["results"])

    def status_for(c):
        return "рекомендовано" if c >= THR_OK else "спірно" if c >= THR_DISPUTED else "окрема сторінка не потрібна"

    def decide(page_counts, kwreg, topic, kwfam):
        fam = collections.Counter()
        for pg, c in page_counts.items():
            if pg:
                fam[R.family(pg)] += c
        if not fam:
            return dict(page=R.parent_page(None, kwreg), cand=None, count=0, family=None, status="окрема сторінка не потрібна", note="", ranked=[])
        f, c = sorted(fam.items(), key=lambda x: (-x[1], R.FAM_PRIORITY.index(x[0])))[0]
        cand, count = None, c
        if f in R.FAM_COUNTRY_PAGE:
            cand = R.FAM_COUNTRY_PAGE[f]
            if kwreg:
                rp = f"{PB}/{kwreg}/{R.FAM_SUFFIX[f]}"
                if page_counts.get(rp, 0) >= THR_DISPUTED:
                    cand, count = rp, page_counts[rp]
        elif f == "hub":
            cand = PB
            if kwreg and page_counts.get(f"{PB}/{kwreg}", 0) >= THR_DISPUTED:
                cand, count = f"{PB}/{kwreg}", page_counts[f"{PB}/{kwreg}"]
        elif f == "hotel":
            cand = R.hotels_page
        elif f == "info":
            cand = f"{PB}/{topic}"
        else:
            cand = R.ROOT_PAGE_FOR.get(f, PB)
        status = status_for(count)
        page, note = cand, ""
        if status == "спірно":
            v = VERDICTS.get(cand)
            if v and kwfam is not None and kwfam != f:
                v = None
            if v:
                if v[0] == "фільтр":
                    page, status = R.parent_page(cand, kwreg), "спірно → фільтр, не окрема сторінка"
                else:
                    status = "спірно → підтверджено перевіркою"
                note = v[1]
            else:
                page = R.FAM_COUNTRY_PAGE[f] if (f in R.FAM_COUNTRY_PAGE and c >= THR_OK and cand != R.FAM_COUNTRY_PAGE[f]) else R.parent_page(cand, kwreg)
                status = "спірно (не перевірено)"
        elif status == "окрема сторінка не потрібна":
            page = R.parent_page(cand, kwreg)
        ranked = sorted(((pg, n) for pg, n in page_counts.items() if pg), key=lambda x: (-x[1], len(x[0])))[:3]
        return dict(page=page, cand=cand, count=count, family=f, fam_count=c, status=status, note=note, ranked=ranked)

    dec = {k: decide(kw_pages[k], R.detect_region(k.lower()), R.topic_for(k), R.keyword_family(k)) for k in order}
    for k, why in MANUAL.items():
        if k in dec:
            dec[k]["status"], dec[k]["note"] = "перевірити вручну", why
    for k, (opage, why) in KEYWORD_OVERRIDES.items():            # owner decisions
        if k in dec:
            dec[k]["serp_page"], dec[k]["serp_status"] = dec[k]["page"], dec[k]["status"]
            dec[k]["page"], dec[k]["status"], dec[k]["note"] = opage, "за рішенням власника", f"{why} За ТОП було б: {dec[k]['serp_page']} ({dec[k]['serp_status']})."
    for k in order:                                               # business rules
        for rule in RULES:
            if k in KEYWORD_OVERRIDES or k in rule.get("exceptions", []):
                continue
            kl = k.lower()
            if (re.search(rule["match"], kl) and not (rule.get("exclude_region") and R.detect_region(kl))
                    and not any(re.search(x, kl) for x in rule.get("exclude_patterns", []))):
                d = dec[k]
                pc = sum(c for pg, c in kw_pages[k].items() if R.family(pg) == rule["family"])
                if d["page"] != rule["page"] or pc < THR_OK:
                    d["serp_page"], d["serp_status"] = d["page"], d["status"]
                    d["note"] = (f"{rule.get('note', '')} У ТОП {rule['family']}-сторінок {pc}/{kw_n[k]}; за ТОП було б: {d['page']} ({d['status']})."
                                 + (f" {d['note']}" if d.get("note") else ""))
                    d["page"], d["status"], d["business"] = rule["page"], "за бізнес-правилом", rule.get("name", "")

    groups = collections.OrderedDict()
    for k in order:
        groups.setdefault(R.lemma_key(k), []).append(k)
    base = sorted(((sum(K[m]["volume"] for m in members), key, members) for key, members in groups.items()), key=lambda x: -x[0])
    H = []
    for bi, (_, key, members) in enumerate(base, 1):
        by_page = collections.OrderedDict()
        for m in members:
            by_page.setdefault(dec[m]["page"], []).append(m)
        parts = sorted(by_page.items(), key=lambda kv: -sum(K[m]["volume"] for m in kv[1]))
        for pi, (page, mem) in enumerate(parts):
            suffix = "" if len(parts) == 1 else "abcdefgh"[pi]
            name, intent_ov = R.CLUSTERS[key]
            it = collections.Counter()
            for m in mem:
                it[K[m]["intent"]] += K[m]["volume"]
            tc, n = collections.Counter(), 0
            for m in mem:
                tc.update(kw_types[m])
                n += kw_n[m]
            toplab, topc = tc.most_common(1)[0]
            H.append(dict(id=f"K{bi:02d}{suffix}", key=key, name=name, members=mem, head=mem[0], total=sum(K[m]["volume"] for m in mem),
                          intent=intent_ov or it.most_common(1)[0][0], page=page, top_type=toplab, top_share=topc / n if n else 0,
                          top_count=f"{topc}/{n}", statuses=dict(collections.Counter(dec[m]["status"] for m in mem)), split=len(parts) > 1, base=f"K{bi:02d}"))
    H.sort(key=lambda c: -c["total"])
    for c in H:
        c["label"] = f"{c['id']} · {c['head']}"

    # ---- SERP overlaps: reference only ----
    keys = list(Sset)
    dist = collections.Counter(len(Sset[x] & Sset[y]) for i, x in enumerate(keys) for y in keys[i + 1:])
    hard, soft = serp_cluster(order, Sset, hard=4), serp_cluster(order, Sset, soft=3)
    cross = []
    for i in range(len(H)):
        for j in range(i + 1, len(H)):
            best = (0, None, None)
            for a in H[i]["members"]:
                for b in H[j]["members"]:
                    n = len(Sset[a] & Sset[b])
                    if n > best[0]:
                        best = (n, a, b)
            if best[0] >= 3:
                same = H[i]["page"] == H[j]["page"]
                cross.append([H[i]["label"], H[j]["label"], best[0], f"{best[1]}  ↔  {best[2]}",
                              "обидва кластери на одній сторінці" if same else
                              f"ПЕРЕТИН ВИДАЧІ: ≥3 спільних URL, а сторінки різні ({H[i]['page']} і {H[j]['page']}): ризик канібалізації, розвести інтент"])
    cross.sort(key=lambda r: -r[2])

    # ---- region extensions (keywords from extend_region/add_region; decisions of the owner, SERP only for some) ----
    EK = {k["keyword"]: k for k in kwj["keywords"] if is_ext(k)}
    reg_serp = {}
    for f in glob.glob(os.path.join(workdir, "*-demand.json")):
        reg_serp.update(json.load(open(f, encoding="utf-8")).get("serp", {}))
    K.update(EK)
    regions_ext = collections.OrderedDict()
    for k, x in EK.items():
        regions_ext.setdefault(ext_region(x), []).append(k)
    ext_hotel_names = {}
    for slug, kws in regions_ext.items():
        rpage = DEC["region_pages"].get(slug) or f"{PB}/{slug}"
        for k in kws:
            x = EK[k]
            tc = collections.Counter(r["type"] for r in reg_serp.get(k, {}).get("results", []))
            kw_types[k], kw_n[k] = tc, sum(tc.values())
            lm = x.get("page_override")
            grp = ext_group(x)
            dec[k] = dict(page=lm or rpage, cand=None, count=0, family=None, status="за рішенням власника", ranked=[],
                          note=(f"Матриця {slug} + тип туру (рішення власника)" if lm else f"Сторінка {slug}: " + {"general": "опис, відпочинок, пляж, FAQ", "hotels_list": "список готелів з фільтрами"}.get(grp, "")))
        letter = slug[:1].upper()
        names = collections.OrderedDict([("general", f"{slug.title()}: загальні запити (опис, відпочинок, FAQ)"), ("hotels_list", f"{slug.title()}: підбірки готелів"),
                                         ("matrix", f"{slug.title()} + тип туру (матриця)")])
        n_id = 0
        for grp, nm in names.items():
            mem = sorted([k for k in kws if ext_group(EK[k]) != "hotel_name" and ("matrix" if EK[k].get("page_override") else ext_group(EK[k])) == grp], key=lambda x: -K[x]["volume"])
            if not mem:
                continue
            n_id += 1
            it = collections.Counter()
            for m in mem:
                it[K[m]["intent"]] += K[m]["volume"]
            tc, n = collections.Counter(), 0
            for m in mem:
                tc.update(kw_types[m])
                n += kw_n[m]
            toplab, topc = tc.most_common(1)[0] if tc else ("н/д (SERP не збирався)", 0)
            lid = f"{letter}{n_id:02d}"
            H.append(dict(id=lid, key=f"{slug}_{grp}", name=nm, members=mem, head=mem[0], total=sum(K[m]["volume"] for m in mem), intent=it.most_common(1)[0][0],
                          page=dec[mem[0]]["page"], top_type=toplab, top_share=(topc / n if n else 0), top_count=(f"{topc}/{n}" if n else "—"),
                          statuses={dec[mem[0]]["status"]: len(mem)}, split=False, base=lid, label=f"{lid} · {mem[0]}"))
        ext_hotel_names[slug] = sorted([k for k in kws if ext_group(EK[k]) == "hotel_name"], key=lambda x: -K[x]["volume"])
    pages = collections.OrderedDict()
    for c in sorted(H, key=lambda c: -c["total"]):
        pages.setdefault(c["page"], []).append(c)

    # ---- serp-data.json ----
    sd = {"meta": {"source": "DataForSEO /v3/serp/google/organic/live/regular", "request_params": next(iter(serp.values()))["request"], "keywords_count": len(serp),
                   "note": "Лише органічні результати, position = rank_group; type/region/page — класифікація посадкових (analyze.py)."},
          "keywords": {k: {"volume": K[k]["volume"], "fetched_at": serp[k]["fetched_at"], "api_datetime": serp[k]["api_datetime"], "check_url": serp[k]["check_url"],
                           "results": serp[k]["results"]} for k in order}}
    json.dump(sd, open(os.path.join(out, "serp-data.json"), "w", encoding="utf-8"), ensure_ascii=False, indent=1)

    # ---- xlsx ----
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook()

    def sheet(ws, head, rows, widths, wrap_cols=()):
        ws.append(head)
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="1F4E78")
            c.alignment = Alignment(vertical="center", wrap_text=True)
        for r in rows:
            ws.append(r)
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        for wc in wrap_cols:
            for r in ws.iter_rows(min_row=2):
                r[wc].alignment = Alignment(wrap_text=True, vertical="top")
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

    def ownpos(k):
        if k not in serp:
            return "не перевірялось"
        site_dom = C.reg_domain(re.sub(r"^https?://", "", BASE))
        p = [r["position"] for r in serp[k]["results"] if C.reg_domain(r["domain"]) == site_dom]
        return min(p) if p else LABS_POS.get(k, "поза ТОП-10")

    def page_state(p):
        return "існує" if p in EXISTING else "нова"

    ws0 = wb.active
    ws0.title = "Увага"
    ws0.column_dimensions["A"].width = 150
    ws0.append(["ЯК ВИЗНАЧЕНО СТОРІНКУ"])
    ws0["A1"].font = Font(bold=True, size=14)
    note_txt = (f"Дані SERP: DataForSEO live/regular, location_code {S.location_code} + se_domain {S.se_domain} (папка {args.serp_dir}). "
                "Сторінку для ключа визначає тип посадкових сторінок конкурентів у ТОП (правила: " + os.path.basename(rules_path) + "). "
                f"Поріг: ≥{THR_OK} — рекомендовано, {THR_DISPUTED}–{THR_OK - 1} — спірно, ≤{THR_DISPUTED - 1} — окрема сторінка не потрібна. "
                "Кластери за лемою/інтентом — основа групування; розщеплюються лише там, де ключі кластера ведуть на різні сторінки. "
                "SERP-кластеризація hard-4/soft-3 і перетини — лише довідка. Ручні рішення зберігаються у semantics-decisions.json; "
                "мінімального порогу обсягу для окремої сторінки немає. Переклади (укр.) — у keywords.json (translation_uk).")
    ws0.append([note_txt])
    ws0["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    ws0.row_dimensions[2].height = 120

    ws = wb.create_sheet("Ключі")
    rows = []
    for c in H:
        for k in sorted(c["members"], key=lambda x: -K[x]["volume"]):
            d = dec[k]
            tl, tcount = kw_types[k].most_common(1)[0] if kw_types[k] else ("н/д (SERP не збирався)", 0)
            extra = (f"найкращий кандидат: {d['cand']} ({d['count']}/{kw_n[k]})" if d.get("cand") else "") + (f". {d['note']}" if d.get("note") else "")
            rows.append([c["label"], k, K[k].get("translation_uk", ""), K[k].get("translation_note_uk", ""), K[k]["volume"],
                         K[k].get("kd") if K[k].get("kd") is not None else "н/д", K[k]["cpc"], ITR[K[k]["intent"]], ownpos(k), BASE + c["page"], tl,
                         (f"{tcount}/{kw_n[k]}" if kw_n[k] else "—"), d["status"], extra])
    sheet(ws, ["кластер", "ключ", "переклад (укр.)", "примітка до перекладу", "частотність", "KD", "CPC (€)", "інтент", f"позиція {BASE.split('//')[-1]}", "рекомендована сторінка",
               "тип сторінки в ТОП", "частка", "статус рішення", "примітка"], rows, [46, 52, 48, 52, 13, 8, 10, 16, 34, 52, 26, 9, 32, 80], wrap_cols=(3, 13))
    ws2 = wb.create_sheet("Кластери")
    rows = []
    for c in H:
        note = NOTES.get(c["page"], "Регіональна сторінка або матриця регіон × тип: рекомендована за типом сторінок конкурентів у ТОП.")
        split = " Лемний кластер розщеплено: частина його ключів веде на іншу сторінку." if c["split"] else ""
        rows.append([c["label"], c["head"], K[c["head"]].get("translation_uk", ""), c["total"], len(c["members"]), ITR[c["intent"]],
                     f"{page_state(c['page']).upper()} → {c['page']}. {note}{split}", c["name"], c["top_type"], c["top_count"], f"{c['top_share']:.0%}",
                     ", ".join(f"{s}: {n}" for s, n in c["statuses"].items())])
    sheet(ws2, ["кластер", "головний ключ", "переклад головного ключа", "сумарна частотність", "к-ть ключів", "інтент", "рекомендація", "назва кластера (лема/інтент)",
                "тип сторінки в ТОП", "к-ть у ТОП", "частка", "статуси рішень по ключах"], rows, [46, 44, 46, 14, 10, 16, 110, 60, 26, 11, 9, 60], wrap_cols=(2, 6, 7, 11))
    ws3 = wb.create_sheet("Типи посадкових конкурентів")
    sheet(ws3, ["ключ", "позиція", "URL", "домен", "тип посадкової", "регіон", "сторінка-кандидат"], URLS, [48, 9, 90, 28, 26, 14, 44])
    labels = sorted({u[4] for u in URLS})
    wsm = wb.create_sheet("Матриця типів по ключах")
    sheet(wsm, ["ключ", "частотність", "органічних"] + labels + ["найкращий кандидат", "кількість", "статус"],
          [[k, K[k]["volume"], kw_n[k]] + [kw_types[k].get(l, 0) for l in labels] + [dec[k].get("cand") or "—", dec[k]["count"], dec[k]["status"]] for k in order],
          [48, 12, 11] + [14] * len(labels) + [44, 10, 32])
    for slug, hn in ext_hotel_names.items():
        if not hn:
            continue
        wsh = wb.create_sheet(("Готелі " + slug.title())[:31])
        rowsh = []
        for kk in hn:
            x = K[kk]
            topt = ", ".join(f"{t} {n}/{kw_n[kk]}" for t, n in kw_types[kk].most_common(2)) if kw_types[kk] else "н/д (SERP не збирався)"
            rowsh.append([kk, x["volume"], x.get("translation_uk", ""), topt, "сторінка окремого готелю на сайті (не сторінка регіону)", x.get("translation_note_uk", "")])
        rowsh.append(["РАЗОМ", sum(K[k]["volume"] for k in hn), "", "", "", ""])
        sheet(wsh, ["ключ", "обсяг", "переклад (укр.)", "тип ТОП", "рекомендація", "примітка"], rowsh, [48, 10, 56, 44, 52, 40])
    # ---- topic boundaries: adjacent directions + long tail (scope.py, profile "scope") ----
    import scope as SC
    SCP = SC.Scope(C.load_profile(S, workdir), R, S)
    SCP.districts = list(DEC["region_pages"])      # districts with own page (e.g. Lara) are separate directions from their parent region
    adj_items = []
    seen_adj = set()
    for k, x in list(K.items()) + list(EK.items()):
        if k in seen_adj or ext_group(x) == "hotel_name":      # ext keywords are also in K; hotel names have their own sheet
            continue
        seen_adj.add(k)
        lab, theme, _ = SCP.label(k)
        if lab == SC.ADJ:
            adj_items.append({"keyword": k, "volume": x["volume"], "_theme": theme, "_now": dec[k]["page"] if k in dec else (x.get("page_override") or DEC["region_pages"].get(ext_region(x)) or f"{PB}/{ext_region(x)}")})
    for x in kwj.get("adjacent", []):
        adj_items.append({"keyword": x["keyword"], "volume": x["volume"], "_theme": x["theme"], "_now": "— (не в списку сторінки)"})
    ws_adj = wb.create_sheet("Напрямки розширення")
    adj_rows = []
    for r in SCP.themes(adj_items):
        nows = collections.Counter(i["_now"] for i in r["items"])
        adj_rows.append([r["theme"], r["n"], r["volume"], r["examples"], SC.recommend(r, SCP.products, EXISTING, set(DEC["region_pages"]), PB),
                         "; ".join(f"{pg} ({n})" for pg, n in nows.most_common(3))])
    adj_rows.append(["РАЗОМ", sum(r[1] for r in adj_rows), sum(r[2] for r in adj_rows), "", "", ""])
    sheet(ws_adj, ["тема (напрямок)", "к-ть ключів", "сумарний обсяг", "приклади ключів (обсяг)", "рекомендація", "де ключі зараз"], adj_rows, [34, 12, 15, 90, 60, 50])
    print(f"Напрямки розширення: {len(adj_rows) - 1} тем, {adj_rows[-1][1]} ключів, обсяг {adj_rows[-1][2]}")
    if kwj.get("longtail"):
        ws_lt = wb.create_sheet("Довгий хвіст")
        sheet(ws_lt, ["ключ", "обсяг", "переклад (укр.)", "інтент", "примітка"],
              [[x["keyword"], x["volume"], x.get("translation_uk", ""), ITR.get(x.get("intent"), "н/д"), "ядро, не вмістилось у ліміт; без кластеризації — для копірайтера"] for x in sorted(kwj["longtail"], key=lambda z: -z["volume"])],
              [55, 10, 50, 16, 60])
    ws4 = wb.create_sheet("Відфільтровані")
    sheet(ws4, ["ключ", "причина"], [[f["keyword"], f["reason"]] for f in kwj.get("filtered", [])], [55, 75])
    ws5 = wb.create_sheet("Конкуренти")
    sheet(ws5, ["ключ", "позиція", "URL", "title з видачі", "опис (сніпет) з видачі", "тип посадкової"],
          [[k, r["position"], r["url"], r["title"], r["snippet"], R.type_label(r["type"], r["region"])] for k in order for r in serp[k]["results"]], [48, 9, 70, 60, 90, 26])
    ws6 = wb.create_sheet("Розподіл по сторінках")
    rows = []
    for p, cs in pages.items():
        tcs, tot_n = collections.Counter(), 0
        for c in cs:
            for m in c["members"]:
                tcs.update(kw_types[m])
                tot_n += kw_n[m]
        topl, topc = tcs.most_common(1)[0]
        v = VERDICTS.get(p)
        why = ("Один кластер — одна сторінка." if len(cs) == 1 else
               f"Кілька кластерів ({', '.join(c['id'] for c in cs)}) ведуть на одну сторінку: для них у ТОП домінує той самий тип посадкових, окремі сторінки під кожен не потрібні.")
        rows.append([BASE + p, page_state(p), len(cs), sum(c["total"] for c in cs), ", ".join(c["id"] for c in cs),
                     f"{topl}: {topc}/{tot_n} ({topc / tot_n:.0%})", f"{v[0]}: {v[1]}" if v else "", why + " " + NOTES.get(p, "")])
    parents = collections.OrderedDict()                # parent pages of matrix pages are always shown, even with 0 own keywords
    for p in pages:
        rel = p[len(PB):].strip("/").split("/") if p.startswith(PB) else []
        if len(rel) == 2 and f"{PB}/{rel[0]}" not in pages:
            parents.setdefault(f"{PB}/{rel[0]}", []).append(p)
    for par, kids in parents.items():
        rows.append([BASE + par, page_state(par), 0, 0, "—", "—", "", "Батьківська сторінка для матриць " + " і ".join("/" + x.rsplit("/", 1)[1] for x in kids) + " (власних ключів немає). " + NOTES.get(par, "")])
    sheet(ws6, ["сторінка", "стан", "кластерів", "сума частотності", "кластери", "домінуючий тип у ТОП", "перевірка спірних", "чому так"], rows, [46, 10, 11, 16, 50, 36, 80, 110], wrap_cols=(6, 7))
    vf = os.path.join(workdir, "landing-verification.json")
    if os.path.exists(vf):
        wsv = wb.create_sheet("Перевірка спірних")
        rowsv = []
        for x in json.load(open(vf, encoding="utf-8")):
            a_, b_ = x["a"], x["b"]
            rowsv.append([x["id"], x["why"], a_.get("url"), a_.get("status"), "; ".join(a_.get("h1") or []), a_.get("words"), a_.get("own_text_words"),
                          b_.get("url"), b_.get("status"), "; ".join(b_.get("h1") or []), b_.get("words"), b_.get("own_text_words"), x.get("text_overlap_6gram")])
        sheet(wsv, ["пара", "що порівнюємо", "URL A", "HTTP A", "H1 A", "слів A", "власний текст A", "URL B", "HTTP B", "H1 B", "слів B", "власний текст B", "збіг тексту (6-грами)"],
              rowsv, [34, 34, 60, 8, 40, 9, 12, 60, 8, 40, 9, 12, 12])
    ws7 = wb.create_sheet("Перетини між кластерами")
    sheet(ws7, ["кластер A", "кластер B", "макс. спільних URL", "ключі з макс. перетином", "висновок"], cross, [46, 46, 14, 70, 110], wrap_cols=(4,))
    sheet(wb.create_sheet("SERP hard-4 (довідка)"), ["кластер", "головний ключ", "сумарна частотність", "к-ть ключів", "ключі"],
          [[f"H{i:02d}", c[0], sum(K[k]["volume"] for k in c), len(c), "; ".join(c)] for i, c in enumerate(sorted(hard, key=lambda c: -sum(K[k]["volume"] for k in c)), 1)], [8, 46, 14, 10, 150])
    sheet(wb.create_sheet("SERP soft-3 (довідка)"), ["кластер", "головний ключ", "сумарна частотність", "к-ть ключів", "ключі"],
          [[f"S{i:02d}", c[0], sum(K[k]["volume"] for k in c), len(c), "; ".join(c)] for i, c in enumerate(sorted(soft, key=lambda c: -sum(K[k]["volume"] for k in c)), 1)], [8, 46, 14, 10, 150])
    target = os.path.join(out, f"semantics-{S.slug}.xlsx")
    try:
        wb.save(target)
    except PermissionError:                                    # file is open in Excel
        target = os.path.join(out, f"semantics-{S.slug}.new.xlsx")
        wb.save(target)
        print("УВАГА: xlsx відкритий в Excel, збережено як", target, "— закрийте Excel і запустіть ще раз.")

    # ---- cluster-map.html ----
    COL = ["#4E79A7", "#F28E2B", "#E15759", "#76B7B2", "#59A14F", "#EDC948", "#B07AA1", "#FF9DA7"]
    SL = R.SLUG
    fam = [("Ця сторінка: " + PB, lambda p: p == PB), ("Pauschalreise", lambda p: p.endswith("/" + SL["pauschal"])), ("Last minute", lambda p: p.endswith("/" + SL["lastminute"])),
           ("Rundreisen", lambda p: p.endswith("/" + SL["rundreise"])), ("Інформаційні", lambda p: p.rsplit("/", 1)[-1] in R.INFO_SLUGS),
           ("All inclusive", lambda p: p.endswith("/" + SL["allinclusive"])), ("Регіони й матриці", lambda p: p.startswith(PB + "/")), ("Готелі", lambda p: p == R.hotels_page)]
    cl = [{"name": n, "color": COL[i], "posts": []} for i, (n, _) in enumerate(fam)]
    for c in H:
        g = next(i for i, (_, f) in enumerate(fam) if f(c["page"]))
        cl[g]["posts"].append({"title": c["head"], "keyword": f"{c['id']} · {c['name']} · {len(c['members'])} ключ. → {c['page']}", "volume": c["total"],
                               "template": page_state(c["page"]), "wordCount": c["total"], "url": c["page"], "status": "planned"})
    keep = [i for i, c in enumerate(cl) if c["posts"]]
    links = []
    for new, old in enumerate(keep):
        for pi in range(len(cl[old]["posts"])):
            nid = f"cluster-{new}-post-{pi}"
            links.append({"from": nid, "to": "pillar", "type": "mandatory"})
            links.append({"from": "pillar", "to": nid, "type": "mandatory" if old == 0 else "recommended"})
    tot = sum(c["total"] for c in H)
    data = {"pillar": {"title": PB, "keyword": main_kw, "volume": K[main_kw]["volume"] if main_kw in K else 0, "template": "поточна сторінка", "wordCount": tot, "url": PB},
            "clusters": [cl[i] for i in keep], "links": links, "meta": {"totalPosts": len(H), "totalClusters": len(keep), "totalLinks": len(links), "estimatedWords": tot}}
    h = open(os.path.join(C.SKILL_DIR, "assets", "cluster-map.html"), encoding="utf-8").read()
    s0, e0 = h.index("const CLUSTER_DATA = {"), h.index("// === END CLUSTER DATA ===")
    h = h[:s0] + "const CLUSTER_DATA = " + json.dumps(data, ensure_ascii=False) + ";\n    " + h[e0:]
    rep = {'lang="en"': 'lang="uk"', "Content Cluster Map": f"Карта кластерів: {S.country}", '"Cluster map for: "': '"Карта кластерів для: "', "Keyword:": "Кластер:", "Volume:": "Сума частотності:",
           "Template:": "Стан сторінки:", "Words:": "Σ частотність:", "Status:": "Статус:", "Est. Words": "Σ частотність", "Mandatory link": "Лінк (ця сторінка ↔ кластер)",
           "Recommended link": "Рекомендований лінк (хаб → окрема сторінка)", "Optional link": "Додатковий лінк", '"Pillar page: "': '"Сторінка: "', '"Spoke page: "': '"Кластер: "',
           "Total Posts": "Кластерів", '<div class="stat-label">Clusters</div>': '<div class="stat-label">Груп (сторінок)</div>', "Internal Links": "Лінків",
           "</div> Written": "</div> Написано", "</div> Planned": "</div> Рекомендовано", "Generated by Claude SEO": "semantics-travel · сторінки за типом посадкових конкурентів"}
    for a_, b_ in rep.items():
        h = h.replace(a_, b_)
    open(os.path.join(out, "cluster-map.html"), "w", encoding="utf-8").write(h)
    json.dump({"clusters": H, "cross": cross, "decisions": dec, "types": {k: dict(v) for k, v in kw_types.items()}},
              open(os.path.join(out, "clusters.json"), "w", encoding="utf-8"), ensure_ascii=False, indent=1)

    print(f"ключів у роботі {len(order)} (+ {len(EK)} ключів розширення регіонів); кластерів {len(H)} (лемних {len(base)}); пар без спільних URL {dist[0] / max(1, sum(dist.values())):.1%}; SERP hard-4 {len(hard)}, soft-3 {len(soft)}")
    for p, cs in pages.items():
        print(f"{p:<42} {page_state(p):<7} clusters {','.join(c['id'] for c in cs)} vol {sum(c['total'] for c in cs)}")
    print("статуси:", dict(collections.Counter(d['status'] for d in dec.values())))
    print("рівень рішень: ручні рішення — у semantics-decisions.json; правила —", rules_path)
    import datetime   # journal of the direction: only in the working folder's CLAUDE.md
    C.journal_add(workdir, "semantics-travel", "analyze — ", f"- analyze — {datetime.date.today().isoformat()}; semantics-{S.slug}.xlsx; ключів {len(order)}, кластерів {len(H)}, сторінок {len(pages)}")


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--workdir", default=".")
    ap.add_argument("--claude-md")
    ap.add_argument("--out-dir")
    ap.add_argument("--serp-dir", default="serp-raw-regular")
    args = ap.parse_args()
    try:
        S = C.load_settings(args.workdir, args.claude_md)
    except C.ConfigMissing as e:
        sys.exit(C.config_help(e))
    run(os.path.abspath(args.workdir), S, args)


if __name__ == "__main__":
    main()
