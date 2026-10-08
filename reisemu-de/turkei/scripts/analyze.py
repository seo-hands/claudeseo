#!/usr/bin/env python3
"""Lemma/intent clusters + page choice by the TYPE OF COMPETITOR LANDING PAGES in the TOP (reusable).

Input : keywords.json and serp-raw-regular/*.json (raw DataForSEO live/regular responses written by
        scripts/fetch_serp.py).  NO API calls are made here.
Output: serp-data.json, semantics-turkei.xlsx, cluster-map.html, clusters.json

python scripts/analyze.py [--serp-dir serp-raw-regular]

How a page is chosen (all rules are editable in the EDITABLE blocks below)
  1. Every organic URL of a keyword's TOP is classified (classify_url): country hub / region hub /
     last minute / all inclusive / pauschalreise (country or region) / hotel / info / rundreise /
     Frühbucher / other.  Each type maps to a candidate landing page (page_for).
  2. Count of the best candidate page among the organic results of the keyword:
        >= THR_OK (5)   -> recommend a page of that type
        THR_DISPUTED..4 -> "спірно": check (landing-verification.json, CONTESTED_VERDICTS)
        <= 2            -> no separate page: keyword goes to the broader (parent) page
  3. Lemma clusters stay the base of grouping; a lemma cluster is split only where its keywords
     land on different pages (K07 -> K07a, K07b ...).
"""
import argparse, collections, glob, json, os, re, sys
from urllib.parse import urlsplit, unquote

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TEMPLATE = os.path.expanduser(r"~\.claude\plugins\cache\agricidaniel-claude-seo\claude-seo\2.4.2\skills\seo-cluster\templates\cluster-map.html")
BASE = "https://reisemu.de"
ITR = {"commercial": "комерційний", "transactional": "транзакційний", "informational": "інформаційний",
       "navigational": "навігаційний", None: "н/д"}
LABS_POS = {"last minute türkei": "23 (Labs, за «türkei last minute»)",
            "last minute urlaub türkei all inclusive": "23 (Labs)"}
EXISTING = {"/tour/turkei", "/tour/turkei/all-inclusive", "/tour/turkei/side", "/tour/turkei/antalya"}  # pages confirmed to exist

# =====================================================================================
# EDITABLE 1: lemma clusters (key -> name, intent override).  Pages are NOT set here.
# =====================================================================================
CLUSTERS = {
    "urlaub_core": ("Türkei Urlaub: бронювання, рік, з перельотом", None),
    "urlaub_price": ("Türkei Urlaub: günstig / billig / Schnäppchen", None),
    "urlaub_season": ("Türkei Urlaub за місяцями (Mai, September, Oktober, November, Dezember)", None),
    "reise": ("Türkei Reise / Reisen: buchen, nach", None),
    "allinclusive": ("Türkei All Inclusive", None),
    "pauschal": ("Pauschalreise Türkei", None),
    "lastminute": ("Last Minute Türkei", None),
    "side": ("Side (Türkische Riviera)", None),
    "antalya": ("Antalya / Lara", None),
    "rundreise_core": ("Türkei Rundreise: загальні, ціна, на власний розсуд", None),
    "rundreise_dauer": ("Türkei Rundreise за тривалістю (7/10/14 днів, 2 тижні)", None),
    "rundreise_region": ("Türkei Rundreise: Kappadokien, Istanbul, Schwarzmeer", None),
    "rundreise_baden": ("Türkei Rundreise + Badeurlaub", None),
    "hotel": ("Hotels Türkei", None),
    "info_warnung": ("Türkei: Reisewarnung, aktuelle Lage, Sicherheit", "informational"),
    "info_einreise": ("Türkei: Einreise mit Personalausweis", "informational"),
}


def lemma_key(k):
    """Lemma/intent cluster of a keyword (first matching rule wins; edit freely)."""
    t = k.lower().replace("-", " ")
    if re.search(r"\bside\b", t):
        return "side"
    if re.search(r"antalya|\blara\b", t):
        return "antalya"
    if re.search(r"warnung|aktuell|sicher", t):
        return "info_warnung"
    if "personalausweis" in t:
        return "info_einreise"
    if t.startswith("hotel"):
        return "hotel"
    if "rundreise" in t:
        if re.search(r"baden|badeurlaub", t):
            return "rundreise_baden"
        if re.search(r"\d+ tage|wochen", t):
            return "rundreise_dauer"
        if re.search(r"kappadokien|istanbul|schwarzmeer", t):
            return "rundreise_region"
        return "rundreise_core"
    if re.search(r"last ?minute", t):
        return "lastminute"
    if "pauschal" in t:
        return "pauschal"
    if re.search(r"all ?inclusive|all inklusiv|alles inklusive", t):
        return "allinclusive"
    if re.search(r"september|oktober|november|dezember|\bmai\b", t):
        return "urlaub_season"
    if re.search(r"günstig|billig|schnäppchen", t):
        return "urlaub_price"
    if re.search(r"reise", t) and not re.search(r"urlaub|flug", t):
        return "reise"
    return "urlaub_core"


# =====================================================================================
# EDITABLE 2: classification of competitor landing pages
# =====================================================================================
THR_OK, THR_DISPUTED = 5, 3          # >=5 recommend; 3-4 disputed; <=2 no separate page
HOTEL_TO_REGION_HUB = True           # hotel pages of a region (e.g. Lara hotels) support the REGION page, not a hotels hub
# canonical region -> regex over URL path (hotel pages: path + title).  First match wins.
REGIONS = collections.OrderedDict([
    ("antalya", r"(?<![a-z])(antalya|lara|kundu|belek)(?![a-z])"),
    ("side", r"(?<![a-z])(side|manavgat|colakli)(?![a-z])"),
    ("alanya", r"(?<![a-z])(alanya|avsallar|konakli|okurcalar|mahmutlar|incekum)(?![a-z])"),
    ("kemer", r"(?<![a-z])(kemer|tekirova|goynuk)(?![a-z])"),
    ("bodrum", r"(?<![a-z])(bodrum|turgutreis|gumbet)(?![a-z])"),
    ("marmaris", r"(?<![a-z])(marmaris|icmeler)(?![a-z])"),
    ("fethiye", r"(?<![a-z])(fethiye|oeluedeniz|oludeniz|calis)(?![a-z])"),
    ("istanbul", r"(?<![a-z])istanbul(?![a-z])"),
    ("kappadokien", r"kappadok|cappadoc"),
    ("didim", r"(?<![a-z])didim(?![a-z])"),
    ("kusadasi", r"(?<![a-z])kusadasi(?![a-z])"),
    ("izmir", r"(?<![a-z])izmir(?![a-z])"),
    ("schwarzmeer", r"schwarzmeer"),
])
INFO_DOMAINS = (r"reddit\.|facebook\.|youtube\.|tiktok\.|instagram\.|pinterest\.|wikipedia\.|auswaertiges-amt|diplo\.de|urlaubspiraten|tagesschau|spiegel\.de|focus\.de|bild\.de|stern\.de|merkur\.de|n-tv\.|chip\.de|"
                r"gutefrage\.|reisefrage\.|forum\.|urbia\.|juraforum|faz\.net|zeit\.de|taz\.de|dw\.com|rtl\.de|deutschlandfunk|bundesregierung|wko\.at|mfa\.gov|eda\.admin|ergo\.de|adac\.de/reise|hna\.de|airliners|iww\.de|reisevor9|swp\.de|dtj-online|znaki\.fm|kleinanzeigen|airbnb|ris\.bka")
# domains that sell tours/Rundreisen only
RUNDREISE_DOMAINS = r"tourradar|marco-polo-reisen|studienreisen|gebeco|ijo-reisen|trendtours|traveloptimizer|rundreisen\.de|dimsumreisen|berge-meer|vivido|desired\.de"
# (type, regex over URL path) applied in this order; then title rules; then region/country fallbacks.
PATH_RULES = [
    ("info", r"warnung|sicherheit|einreise|reisehinweis|reisetipps|ratgeber|magazin|/blog|forum|/foren/|reisebericht|inspiration|/wetter|klima|reisezeit|/faq|visum|personalausweis|reisepass|/news|reise-glossar|/thema/"),
    ("lastminute", r"last-?minute|lastminute"),
    ("pauschal", r"pauschal"),
    ("allinclusive", r"all-?inclusive|allinclusive|all-inklusive"),
    ("rundreise", r"rundreise"),
    ("fruehbucher", r"fruehbucher|frühbucher|fruhbucher"),
    ("hotel", r"/(hotel|hotels|hi|ho|hrd|dh|beschreibung|hotelbewertung|hotelbewertungen|hotelangebote|unterkunft|resorts?)(/|$)|-\d{5,}$|_hid-\d+|hotel-resort|hotelslist|[-/]hotels?([-/]|$)"),
]
TITLE_RULES = [
    ("lastminute", r"^(super[ -])?last[ -]?minute"),
    ("pauschal", r"^pauschalreisen?"),
    ("allinclusive", r"^(türkei )?all[ -]?inclusive"),
    ("rundreise", r"rundreis|studienreise|kombi-?reise|\d+-tage-route"),
    ("hotel", r"(hotel|hotels|resort|suites?|palace|beach club|residence|boutique)"),
]
HUB_COUNTRY = re.compile(r"/(urlaub|reisen|reiseziele|urlaubsziele|badereisen|reiseland|laender|europa|land|ferienreisen|strandurlaub|ferien)(/[a-z0-9-]+)*/(tuerkei|turkei|türkei|tuerkei-co4|ferienreisen-tuerkei|tr_tuerkei)(/\d+)?(\.html|\.php)?/?$"
                         r"|/(tuerkei|turkei|türkei|ferienreisen-tuerkei|tuerkei-urlaub|tr_tuerkei)(/tuerkei-urlaub)?(\.html|\.php)?/?$|riviera|aegaeis|region/?$|/urlaub/tuerkei/(inland|[a-z-]*region)|tuerkei-urlaub-buchen|tuerkei/billigurlaub")


def detect_region(text):
    for name, pat in REGIONS.items():
        if re.search(pat, text):
            return name
    return None


def classify_url(url, title=""):
    """-> (type, region).  Types: hub_country, hub_region, lastminute, allinclusive, pauschal, hotel,
    info, rundreise, fruehbucher, other.  `region` is a canonical region name or None."""
    u = urlsplit(url)
    host, path, t = u.netloc.lower(), unquote(u.path).lower(), title.lower()
    if re.search(INFO_DOMAINS, host + path):
        return "info", None
    if re.search(RUNDREISE_DOMAINS, host):
        return "rundreise", None
    region = detect_region(path)
    typ = next((name for name, pat in PATH_RULES if re.search(pat, path)), None)
    if typ is None:
        typ = next((name for name, pat in TITLE_RULES if re.search(pat, t)), None)
        if typ == "hotel" and len([s for s in path.split("/") if s]) < 3:
            typ = None
    if typ == "hotel":
        region = detect_region(path + " " + t)
    if typ:
        return typ, region
    if region:
        return "hub_region", region
    if HUB_COUNTRY.search(path):
        return "hub_country", None
    return "other", None


TYPE_LABEL = {
    ("hub_country", False): "хаб країни", ("hub_region", True): "хаб регіону",
    ("lastminute", False): "last minute (країна)", ("lastminute", True): "last minute (регіон)",
    ("allinclusive", False): "all inclusive (країна)", ("allinclusive", True): "all inclusive (регіон)",
    ("pauschal", False): "pauschalreise (країна)", ("pauschal", True): "pauschalreise (регіон)",
    ("hotel", False): "готель", ("hotel", True): "готель (регіон)", ("info", False): "інформаційна",
    ("rundreise", False): "rundreise", ("fruehbucher", False): "Frühbucher", ("other", False): "інше",
}


def type_label(typ, region):
    return TYPE_LABEL.get((typ, bool(region))) or TYPE_LABEL.get((typ, False)) or typ


def page_for(typ, region, kwreg=None, topic="reisehinweise"):
    """Candidate landing page of the site for a competitor page type (None = no page).
    kwreg = region named in the KEYWORD (hotel pages of that region support the region page);
    topic = guide page for informational results (hinweise | reisezeit | ratgeber)."""
    r = f"/tour/turkei/{region}" if region else None
    if typ == "hub_country":
        return "/tour/turkei"
    if typ == "hub_region":
        return r
    if typ == "lastminute":
        return (r + "/last-minute") if r else "/tour/turkei/last-minute"
    if typ == "fruehbucher":
        return "/tour/turkei/last-minute"          # user rule: last-minute / Frühbucher pages count together
    if typ == "allinclusive":
        return (r + "/all-inclusive") if r else "/tour/turkei/all-inclusive"
    if typ == "pauschal":
        return (r + "/pauschalreise") if r else "/tour/turkei/pauschalreise"
    if typ == "hotel":
        return r if (r and HOTEL_TO_REGION_HUB and region == kwreg) else "/hotels/turkei"
    if typ == "info":
        return f"/tour/turkei/{topic}"
    if typ == "rundreise":
        return "/tour/turkei/rundreisen"
    return None


FAM_BY_LAST = {"last-minute": "lastminute", "all-inclusive": "allinclusive", "pauschalreise": "pauschal", "rundreisen": "rundreise",
               "reisehinweise": "info", "reisezeit": "info", "ratgeber": "info"}
FAM_PRIORITY = ["hub", "lastminute", "allinclusive", "pauschal", "info", "hotel", "rundreise"]
FAM_COUNTRY_PAGE = {"lastminute": "/tour/turkei/last-minute", "allinclusive": "/tour/turkei/all-inclusive", "pauschal": "/tour/turkei/pauschalreise"}
FAM_SUFFIX = {"lastminute": "last-minute", "allinclusive": "all-inclusive", "pauschal": "pauschalreise"}


def family(page):
    """Page family: types of the same product are counted together (country page + region pages)."""
    if page == "/hotels/turkei":
        return "hotel"
    return FAM_BY_LAST.get(page.rsplit("/", 1)[1], "hub")


def parent_page(page, kwreg=None):
    """Broader page when no separate page is justified."""
    if page == "/tour/turkei":
        return page
    return f"/tour/turkei/{kwreg}" if kwreg else "/tour/turkei"


# verdicts for disputed (3-4 of 10) pages after checking competitor pages (see landing-verification.json)
# page -> ("окрема" | "фільтр", reason)
CONTESTED_VERDICTS = {
    "/tour/turkei/all-inclusive": ("окрема", "sonnenklar.tv (/urlaub vs /all-inclusive) і lidl-reisen.de: різні H1, 590–1045 слів власного тексту, збіг тексту ≈0 → окрема посадкова, не фільтр"),
    "/tour/turkei/pauschalreise": ("окрема", "sonnenklar.tv і coraltravel.de: різні H1 («Pauschalreise Türkei», «Türkei Urlaub 2026 – Pauschalreisen…»), 379–1875 слів власного тексту, збіг ≈0 → окрема посадкова"),
    "/tour/turkei/antalya/last-minute": ("окрема", "lidl-reisen.de має обидві сторінки («Urlaub Antalya» і «Last Minute Urlaub Antalya»), ~1000 слів власного тексту кожна, збіг 0; sonnenklar, TUI, ANEX мають окремі Last-Minute-Antalya (200–667 слів). Підтверджено лише частково: пару «хаб + last minute» разом має тільки lidl"),
    "/hotels/turkei": ("окрема", "tui.com/hotels/tuerkei — окрема сторінка («Türkei Hotel», ~780 слів власного тексту) поруч із хабом /urlaub/tuerkei; holidaycheck недоступний (400), restplatzboerse рендериться JS. Застосовано лише до ключів про готелі"),
}
# keywords that must be checked by hand regardless of the rule: keyword -> reason
MANUAL = {"urlaub türkei buchen": "ТОП розкиданий: готелі, регіони, мало спільних URL із «türkei urlaub»; перевірити вручну"}

# Business rules override the SERP-based page.  The keyword is marked "за бізнес-правилом" when the SERP share of
# that page type is below THR_OK (or the SERP would have chosen another page).
# A rule does NOT apply to keywords that name a region (Side, Antalya, Lara ...) or contain another tour type:
#   exclude_patterns - regexes over the keyword (e.g. last minute; all inclusive written BEFORE the rule's own term
#   means all inclusive is the main type), exclude_region - skip keywords with a region name.
BUSINESS_RULES = [
    {"name": "pauschalreise", "match": r"pauschalreisen?", "family": "pauschal", "page": "/tour/turkei/pauschalreise",
     "exclude_region": True,
     "exclude_patterns": [r"last ?minute", r"all[ -]?inclusive.*pauschal"],
     "note": "Хаб не оптимізуємо під Pauschalreise: ключі з «Pauschalreise/Pauschalreisen» без регіону й без іншого головного типу туру ведуть на /tour/turkei/pauschalreise."},
]
# explicit per-keyword decisions of the owner (override both SERP and rules): keyword -> (page, reason)
LARA_PAGE = "/tour/turkei/antalya/lara"
KEYWORD_OVERRIDES = {
    "pauschalreise türkei lara": (LARA_PAGE, "За рішенням власника: окрема сторінка Lara."),
    "pauschalreise türkei side all inclusive": ("/tour/turkei/side/pauschalreise", "За вказівкою: матриця Side + Pauschalreise разом з «pauschalreise side türkei»; All inclusive тут лише уточнення."),
}

PAGE_NOTES = {
    "/tour/turkei": "Хаб країни «Türkei Urlaub»: у ТОП домінують хаби країни.",
    "/tour/turkei/all-inclusive": "⚠ КАНІБАЛІЗАЦІЯ: сторінка існує, але показує той самий блок «Last-Minute-Angebote», що й /tour/turkei.",
    "/tour/turkei/last-minute": "⚠ КАНІБАЛІЗАЦІЯ: /tour/turkei має title «Last-Minute-Angebote» і ранжується за last minute (поз. 23), а ще є загальна /last-minute-reisen.",
    "/tour/turkei/pauschalreise": "Окремий тип сторінки конкурентів (pauschalreise) + бізнес-правило: усі ключі з «Pauschalreise» ведуть сюди, хаб під цей запит не оптимізуємо.",
    "/tour/turkei/reisehinweise": "Інформаційний інтент; низький пріоритет.",
    "/tour/turkei/rundreisen": "Rundreise — інший продукт; сторінки на сайті немає.",
    "/hotels/turkei": "Хаб готелів (готельні сторінки сайту вже є).",
    LARA_PAGE: "Нова сторінка Lara (рішення власника): опис курорту, список готелів з фільтрами 4–5★/All Inclusive, короткий блок «Beste Reisezeit» (без прив'язки до ключів), FAQ. Погодні/картографічні ключі Lara не цільові (відфільтровані: у ТОП погодні сервіси). Окремі готелі (Liberty Lara, Fame Residence…) — на сторінки готелів, див. аркуш «Готелі Lara».",
    "/tour/turkei/reisezeit": "Інформаційний гайд «Beste Reisezeit / Türkei im Monat…»: у ТОП для місяців переважають інформаційні сторінки, а не комерційні.",
    "/tour/turkei/ratgeber": "Інформаційний гайд (Strand, Meer, Tipps): у ТОП переважають інформаційні сторінки.",
}


def nu(u):
    p = urlsplit(u)
    h = p.netloc.lower()
    h = h[4:] if h.startswith("www.") else h
    return h + p.path.lower().rstrip("/")


def load(root, serp_dir="serp-raw-regular"):
    cands = json.load(open(os.path.join(root, "keywords.json"), encoding="utf-8"))
    K = {k["keyword"]: k for k in cands["keywords"] if k.get("source") != "lara"}      # Lara keywords have no SERP of their own (see lara-demand.json)
    serp = {}
    for f in glob.glob(os.path.join(root, serp_dir, "*.json")):
        d = json.load(open(f, encoding="utf-8"))
        res = d["response"]["tasks"][0]["result"][0]
        rows = [{"position": it["rank_group"], "url": it["url"], "domain": it["domain"],
                 "title": it.get("title") or "", "snippet": it.get("description") or ""}
                for it in res["items"] if it["type"] == "organic"]
        serp[d["keyword"]] = {"results": rows, "request": d["request"], "fetched_at": d["fetched_at"],
                              "api_datetime": res.get("datetime"), "check_url": res.get("check_url")}
    missing = [k for k in K if k not in serp]
    if missing:
        sys.exit(f"no raw SERP for: {missing}")
    return cands, K, serp


def serp_cluster(order, S, hard=None, soft=None):
    cl = []
    for k in order:
        best, bs = None, -1
        for c in cl:
            if hard:
                ov = [len(S[k] & S[m]) for m in c]
                ok, score = min(ov) >= hard, sum(ov) / len(ov)
            else:
                ov0 = len(S[k] & S[c[0]])
                ok, score = ov0 >= soft, ov0
            if ok and score > bs:
                best, bs = c, score
        if best is None:
            cl.append([k])
        else:
            best.append(k)
    return cl


def status_for(c):
    return "рекомендовано" if c >= THR_OK else "спірно" if c >= THR_DISPUTED else "окрема сторінка не потрібна"


def keyword_family(k):
    """Product family named in the keyword itself (a verdict about a page type applies only to keywords of that family)."""
    t = k.lower()
    if re.search(r"last ?minute", t):
        return "lastminute"
    if "pauschal" in t:
        return "pauschal"
    if re.search(r"all[ -]?inclusive|all[ -]?inklusiv|alles inklusive", t):
        return "allinclusive"
    if t.startswith("hotel"):
        return "hotel"
    if "rundreise" in t:
        return "rundreise"
    return "hub"


def decide(page_counts, kwreg=None, topic="reisehinweise", lemma=None, kwfam=None):
    """page_counts: Counter candidate-page -> count among the organic results of ONE keyword.
    Counts are summed per product family (country + region pages together).  A region-specific page
    (matrix, e.g. /tour/turkei/antalya/last-minute) is chosen only when the keyword names that region."""
    fam = collections.Counter()
    for pg, c in page_counts.items():
        if pg:
            fam[family(pg)] += c
    if not fam:
        return dict(page=parent_page(None, kwreg), cand=None, count=0, family=None, status="окрема сторінка не потрібна", note="", ranked=[])
    f, c = sorted(fam.items(), key=lambda x: (-x[1], FAM_PRIORITY.index(x[0])))[0]
    cand, count = None, c
    if f in FAM_COUNTRY_PAGE:
        cand = FAM_COUNTRY_PAGE[f]
        if kwreg:
            rp = f"/tour/turkei/{kwreg}/{FAM_SUFFIX[f]}"
            if page_counts.get(rp, 0) >= THR_DISPUTED:
                cand, count = rp, page_counts[rp]
    elif f == "hub":
        cand = "/tour/turkei"
        if kwreg and page_counts.get(f"/tour/turkei/{kwreg}", 0) >= THR_DISPUTED:
            cand, count = f"/tour/turkei/{kwreg}", page_counts[f"/tour/turkei/{kwreg}"]
    elif f == "hotel":
        cand = "/hotels/turkei"
    elif f == "info":
        cand = f"/tour/turkei/{topic}"
    else:
        cand = "/tour/turkei/rundreisen"
    status = status_for(count)
    page, note = cand, ""
    if status == "спірно":
        v = CONTESTED_VERDICTS.get(cand)
        if v and kwfam is not None and kwfam != f:
            v = None          # a verdict about a page type applies only to keywords of that product family
        if v:
            if v[0] == "фільтр":
                page, status = parent_page(cand, kwreg), "спірно → фільтр, не окрема сторінка"
            else:
                status = "спірно → підтверджено перевіркою"
            note = v[1]
        else:
            # matrix page with strong family evidence falls back to the country-level family page
            page = FAM_COUNTRY_PAGE[f] if (f in FAM_COUNTRY_PAGE and c >= THR_OK and cand != FAM_COUNTRY_PAGE[f]) else parent_page(cand, kwreg)
            status = "спірно (не перевірено)"
    elif status == "окрема сторінка не потрібна":
        page = parent_page(cand, kwreg)
    ranked = sorted(((pg, n) for pg, n in page_counts.items() if pg), key=lambda x: (-x[1], len(x[0])))[:3]
    return dict(page=page, cand=cand, count=count, family=f, fam_count=c, status=status, note=note, ranked=ranked)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--serp-dir", default="serp-raw-regular")
    args = ap.parse_args()
    cands, K, serp = load(ROOT, args.serp_dir)
    order = sorted(K, key=lambda k: -K[k]["volume"])
    S = {k: {nu(r["url"]) for r in v["results"]} for k, v in serp.items() if k in K}
    ref = ["schauinsland-reisen", "holidaycheck", "sonnenklar", "anextour", "coraltravel", "aldi-reisen", "restplatzboerse", "check24"]
    hit = sum(1 for r in serp["türkei urlaub"]["results"] if any(x in r["domain"] for x in ref))
    print(f"перевірка «türkei urlaub»: {hit} із {len(ref)} доменів збігаються з ручним ТОП ({args.serp_dir})")

    # ---- classify every URL of every TOP ----
    URLS = []
    kw_types, kw_pages, kw_n = {}, {}, {}
    for k in order:
        tc, pc = collections.Counter(), collections.Counter()
        kwreg = detect_region(k.lower())
        topic = "reisezeit" if lemma_key(k) == "urlaub_season" else "reisehinweise" if lemma_key(k).startswith("info_") else "ratgeber"
        for r in serp[k]["results"]:
            typ, region = classify_url(r["url"], r["title"])
            page = page_for(typ, region, kwreg, topic)
            r["type"], r["region"], r["page"] = typ, region, page
            tc[type_label(typ, region)] += 1
            if page:
                pc[page] += 1
            URLS.append([k, r["position"], r["url"], r["domain"].replace("www.", ""), type_label(typ, region), region or "", page or "—"])
        kw_types[k], kw_pages[k], kw_n[k] = tc, pc, len(serp[k]["results"])

    # ---- keyword-level decision, then lemma cluster x page ----
    dec = {}
    for k in order:
        tp = "reisezeit" if lemma_key(k) == "urlaub_season" else "reisehinweise" if lemma_key(k).startswith("info_") else "ratgeber"
        dec[k] = decide(kw_pages[k], detect_region(k.lower()), tp, lemma_key(k), keyword_family(k))
    for k, why in MANUAL.items():
        if k in dec:
            dec[k]["status"] = "перевірити вручну"
            dec[k]["note"] = why
    for k, (opage, why) in KEYWORD_OVERRIDES.items():       # owner overrides
        if k in dec:
            dec[k]["serp_page"], dec[k]["serp_status"] = dec[k]["page"], dec[k]["status"]
            dec[k]["page"], dec[k]["status"], dec[k]["note"] = opage, "за рішенням власника", f"{why} За ТОП було б: {dec[k]['serp_page']} ({dec[k]['serp_status']})."
    for k in order:                      # business rules
        for rule in BUSINESS_RULES:
            if k in KEYWORD_OVERRIDES:
                continue
            kl = k.lower()
            if (re.search(rule["match"], kl) and not (rule.get("exclude_region") and detect_region(kl))
                    and not any(re.search(x, kl) for x in rule.get("exclude_patterns", []))):
                d = dec[k]
                pc = sum(c for pg, c in kw_pages[k].items() if family(pg) == rule["family"])
                if d["page"] != rule["page"] or pc < THR_OK:
                    d["serp_page"], d["serp_status"] = d["page"], d["status"]
                    d["note"] = (f"{rule['note']} У ТОП {rule['family']}-сторінок {pc}/{kw_n[k]}; за ТОП було б: {d['page']} ({d['status']})."
                                 + (f" {d['note']}" if d.get("note") else ""))
                    d["page"], d["status"], d["business"] = rule["page"], "за бізнес-правилом", rule["name"]
    groups = collections.OrderedDict()
    for k in order:
        groups.setdefault(lemma_key(k), []).append(k)
    base = sorted(((sum(K[m]["volume"] for m in members), key, members) for key, members in groups.items()), key=lambda x: -x[0])
    H = []
    for bi, (_, key, members) in enumerate(base, 1):
        by_page = collections.OrderedDict()
        for m in members:
            by_page.setdefault(dec[m]["page"], []).append(m)
        parts = sorted(by_page.items(), key=lambda kv: -sum(K[m]["volume"] for m in kv[1]))
        for pi, (page, mem) in enumerate(parts):
            suffix = "" if len(parts) == 1 else "abcdefgh"[pi]
            name, intent_ov = CLUSTERS[key]
            it = collections.Counter()
            for m in mem:
                it[K[m]["intent"]] += K[m]["volume"]
            tc, n = collections.Counter(), 0
            for m in mem:
                tc.update(kw_types[m])
                n += kw_n[m]
            toplab, topc = tc.most_common(1)[0]
            sts = collections.Counter(dec[m]["status"] for m in mem)
            H.append(dict(id=f"K{bi:02d}{suffix}", key=key, name=name, members=mem, head=mem[0],
                          total=sum(K[m]["volume"] for m in mem), intent=intent_ov or it.most_common(1)[0][0], page=page,
                          top_type=toplab, top_share=topc / n if n else 0, top_count=f"{topc}/{n}",
                          statuses=dict(sts), split=len(parts) > 1, base=f"K{bi:02d}"))
    H.sort(key=lambda c: -c["total"])
    for c in H:
        c["label"] = f"{c['id']} · {c['head']}"

    # ---- SERP overlap clustering for reference + cross-cluster overlaps ----
    keys = list(S)
    dist = collections.Counter(len(S[x] & S[y]) for i, x in enumerate(keys) for y in keys[i + 1:])
    hard, soft = serp_cluster(order, S, hard=4), serp_cluster(order, S, soft=3)
    cross = []
    for i in range(len(H)):
        for j in range(i + 1, len(H)):
            best = (0, None, None)
            for a in H[i]["members"]:
                for b in H[j]["members"]:
                    n = len(S[a] & S[b])
                    if n > best[0]:
                        best = (n, a, b)
            if best[0] >= 3:
                same = H[i]["page"] == H[j]["page"]
                cross.append([H[i]["label"], H[j]["label"], best[0], f"{best[1]}  ↔  {best[2]}",
                              "обидва кластери на одній сторінці" if same else
                              f"ПЕРЕТИН ВИДАЧІ: ≥3 спільних URL, а сторінки різні ({H[i]['page']} і {H[j]['page']}): ризик канібалізації, розвести інтент"])
    cross.sort(key=lambda r: -r[2])
    # ---- Lara (owner decision: separate page /tour/turkei/antalya/lara; data from lara-demand.json, no SERP of its own) ----
    LK = {k["keyword"]: k for k in cands["keywords"] if k.get("source") == "lara"}
    lf = os.path.join(ROOT, "lara-demand.json")
    lara_serp = json.load(open(lf, encoding="utf-8")).get("serp", {}) if os.path.exists(lf) else {}
    K.update(LK)
    for k, x in LK.items():
        tc = collections.Counter(r["type"] for r in lara_serp.get(k, {}).get("results", []))
        kw_types[k], kw_n[k] = tc, sum(tc.values())
        lm = x.get("page_override")
        dec[k] = dict(page=lm or LARA_PAGE, cand=None, count=0, family=None, status="за рішенням власника",
                      note=("Матриця Antalya + last minute (рішення власника)" if lm else "Сторінка Lara: " + {"general": "опис курорту, відпочинок, пляж, FAQ", "hotels_list": "список готелів з фільтрами 4–5★/All Inclusive"}.get(x["lara_group"], "")), ranked=[])
    lara_names = {"general": ("Lara: відпочинок, курорт, пляж, FAQ", "L01"), "hotels_list": ("Lara: підбірки готелів", "L02"),
                  "lastminute": ("Lara + last minute (матриця Antalya)", "L03")}
    for grp, (nm, lid) in lara_names.items():
        mem = sorted([k for k, x in LK.items() if ("lastminute" if x.get("page_override") else x["lara_group"]) == grp], key=lambda x: -K[x]["volume"])
        if grp == "general" and "pauschalreise türkei lara" in K:
            pass
        it = collections.Counter()
        for m in mem:
            it[K[m]["intent"]] += K[m]["volume"]
        tc, n = collections.Counter(), 0
        for m in mem:
            tc.update(kw_types[m])
            n += kw_n[m]
        toplab, topc = tc.most_common(1)[0] if tc else ("н/д (SERP не збирався)", 0)
        H.append(dict(id=lid, key="lara_" + grp, name=nm, members=mem, head=mem[0], total=sum(K[m]["volume"] for m in mem),
                      intent=it.most_common(1)[0][0], page=dec[mem[0]]["page"], top_type=toplab, top_share=(topc / n if n else 0), top_count=(f"{topc}/{n}" if n else "—"),
                      statuses={dec[mem[0]]["status"]: len(mem)}, split=False, base=lid, label=f"{lid} · {mem[0]}"))
    pages = collections.OrderedDict()
    for c in sorted(H, key=lambda c: -c["total"]):
        pages.setdefault(c["page"], []).append(c)

    # ---- serp-data.json ----
    sd = {"meta": {"source": "DataForSEO API /v3/serp/google/organic/live/regular (scripts/fetch_serp.py)",
                   "request_params": next(iter(serp.values()))["request"], "keywords_count": len(serp),
                   "note": "Лише органічні результати (type=organic), position = rank_group; type/region/page — класифікація посадкових (analyze.py). Сирі відповіді API у serp-raw-regular/."},
          "keywords": {k: {"volume": K[k]["volume"], "fetched_at": serp[k]["fetched_at"], "api_datetime": serp[k]["api_datetime"],
                           "check_url": serp[k]["check_url"], "results": serp[k]["results"]} for k in order}}
    json.dump(sd, open(os.path.join(ROOT, "serp-data.json"), "w", encoding="utf-8"), ensure_ascii=False, indent=1)

    # ---- xlsx ----
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    wb = Workbook()

    def sheet(ws, head, rows, widths, wrap_cols=()):
        ws.append(head)
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="1F4E78")
            c.alignment = Alignment(vertical="center", wrap_text=True)
        for r in rows:
            ws.append(r)
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        for wc in wrap_cols:
            for r in ws.iter_rows(min_row=2):
                r[wc].alignment = Alignment(wrap_text=True, vertical="top")
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

    def ownpos(k):
        if k not in serp:
            return "не перевірялось"
        p = [r["position"] for r in serp[k]["results"] if r["domain"].replace("www.", "") == "reisemu.de"]
        return min(p) if p else LABS_POS.get(k, "поза ТОП-10")

    def page_state(p):
        return "існує" if p in EXISTING else "нова"

    note_txt = ("Дані SERP: DataForSEO live/regular, location_code 2276 + se_domain google.de (2026-10-07, serp-raw-regular/); «türkei urlaub» збігається з ручною перевіркою (8 із 8). "
                "Сторінку для ключа визначає тип посадкових сторінок конкурентів у ТОП (правила в analyze.py). Поріг: ≥5 — рекомендовано, 3–4 — спірно, ≤2 — окрема сторінка не потрібна. "
                "Кластери за лемою/інтентом — основа групування; розщеплюються лише там, де ключі кластера ведуть на різні сторінки. "
                "Бізнес-правило: ключі з «Pauschalreise/Pauschalreisen» (крім «pauschalreise side türkei») ведуть на /tour/turkei/pauschalreise; статус «за бізнес-правилом», якщо SERP-частка <5. "
                "Сторінки /reisezeit, /ratgeber, /side/pauschalreise лишаються окремими, мінімального порогу обсягу немає. Сторінка /antalya/lara — за рішенням власника (ключі Lara з lara-demand.json, назви готелів окремо в аркуші «Готелі Lara»). Переклади (укр.) зберігаються в keywords.json (translation_uk).")
    ws0 = wb.active
    ws0.title = "Увага"
    ws0.column_dimensions["A"].width = 150
    ws0.append(["ЯК ВИЗНАЧЕНО СТОРІНКУ"])
    ws0["A1"].font = Font(bold=True, size=14)
    ws0.append([note_txt])
    ws0["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    ws0.row_dimensions[2].height = 110

    ws = wb.create_sheet("Ключі")
    rows = []
    for c in H:
        for k in sorted(c["members"], key=lambda x: -K[x]["volume"]):
            d = dec[k]
            tl, tcount = kw_types[k].most_common(1)[0] if kw_types[k] else ("н/д (SERP не збирався)", 0)
            extra = (f"найкращий кандидат: {d['cand']} ({d['count']}/{kw_n[k]})" if d.get("cand") else "") + (f". {d['note']}" if d.get("note") else "")
            rows.append([c["label"], k, K[k].get("translation_uk", ""), K[k].get("translation_note_uk", ""), K[k]["volume"], K[k].get("kd") if K[k].get("kd") is not None else "н/д", K[k]["cpc"],
                         ITR[K[k]["intent"]], ownpos(k), BASE + c["page"], tl, (f"{tcount}/{kw_n[k]}" if kw_n[k] else "—"), d["status"], extra])
    sheet(ws, ["кластер", "ключ", "переклад (укр.)", "примітка до перекладу", "частотність", "KD", "CPC (€)", "інтент", "позиція reisemu.de", "рекомендована сторінка",
               "тип сторінки в ТОП", "частка", "статус рішення", "примітка"], rows, [46, 52, 48, 52, 13, 8, 10, 16, 34, 52, 26, 9, 32, 80], wrap_cols=(3, 13))
    ws2 = wb.create_sheet("Кластери")
    rows = []
    for c in H:
        st = ", ".join(f"{s}: {n}" for s, n in c["statuses"].items())
        note = PAGE_NOTES.get(c["page"], "Регіональна сторінка або матриця регіон × тип: рекомендована за типом сторінок конкурентів у ТОП.")
        split = " Лемний кластер розщеплено: частина його ключів веде на іншу сторінку." if c["split"] else ""
        rows.append([c["label"], c["head"], K[c["head"]].get("translation_uk", ""), c["total"], len(c["members"]), ITR[c["intent"]],
                     f"{page_state(c['page']).upper()} → {c['page']}. {note}{split}", c["name"], c["top_type"], c["top_count"], f"{c['top_share']:.0%}", st])
    sheet(ws2, ["кластер", "головний ключ", "переклад головного ключа", "сумарна частотність", "к-ть ключів", "інтент", "рекомендація", "назва кластера (лема/інтент)",
                "тип сторінки в ТОП", "к-ть у ТОП", "частка", "статуси рішень по ключах"], rows, [46, 44, 46, 14, 10, 16, 110, 60, 26, 11, 9, 60], wrap_cols=(2, 6, 7, 11))
    ws3 = wb.create_sheet("Типи посадкових конкурентів")
    sheet(ws3, ["ключ", "позиція", "URL", "домен", "тип посадкової", "регіон", "сторінка-кандидат на reisemu.de"], URLS, [48, 9, 90, 28, 26, 14, 44])
    labels = sorted({u[4] for u in URLS})
    wsm = wb.create_sheet("Матриця типів по ключах")
    sheet(wsm, ["ключ", "частотність", "органічних"] + labels + ["найкращий кандидат", "кількість", "статус"],
          [[k, K[k]["volume"], kw_n[k]] + [kw_types[k].get(l, 0) for l in labels] + [dec[k].get("cand") or "—", dec[k]["count"], dec[k]["status"]] for k in order],
          [48, 12, 11] + [14] * len(labels) + [44, 10, 32])
    wsh = wb.create_sheet("Готелі Lara")
    hn = sorted([x for x in LK.values() if x["lara_group"] == "hotel_name"], key=lambda x: -x["volume"])
    rowsh = []
    for x in hn:
        kk = x["keyword"]
        topt = ", ".join(f"{t} {n}/{kw_n[kk]}" for t, n in kw_types[kk].most_common(2)) if kw_types[kk] else "н/д (SERP не збирався)"
        rowsh.append([kk, x["volume"], x.get("translation_uk", ""), topt, "сторінка окремого готелю на сайті (не сторінка Lara)",
                      x.get("translation_note_uk", "")])
    rowsh.append(["РАЗОМ", sum(x["volume"] for x in hn), "", "", "", ""])
    sheet(wsh, ["ключ", "обсяг", "переклад (укр.)", "тип ТОП", "рекомендація", "примітка"], rowsh, [48, 10, 56, 44, 52, 40])
    ws4 = wb.create_sheet("Відфільтровані")
    sheet(ws4, ["ключ", "причина"], [[f["keyword"], f["reason"]] for f in cands["filtered"]], [55, 75])
    ws5 = wb.create_sheet("Конкуренти")
    sheet(ws5, ["ключ", "позиція", "URL", "title з видачі", "опис (сніпет) з видачі", "тип посадкової"],
          [[k, r["position"], r["url"], r["title"], r["snippet"], type_label(r["type"], r["region"])] for k in order for r in serp[k]["results"]],
          [48, 9, 70, 60, 90, 26])
    ws6 = wb.create_sheet("Розподіл по сторінках")
    rows = []
    for p, cs in pages.items():
        tcs, tot_n = collections.Counter(), 0
        for c in cs:
            for m in c["members"]:
                tcs.update(kw_types[m])
                tot_n += kw_n[m]
        topl, topc = tcs.most_common(1)[0]
        v = CONTESTED_VERDICTS.get(p)
        ver = f"{v[0]}: {v[1]}" if v else ""
        why = ("Один кластер — одна сторінка." if len(cs) == 1 else
               f"Кілька кластерів ({', '.join(c['id'] for c in cs)}) ведуть на одну сторінку: для них у ТОП домінує той самий тип посадкових, окремі сторінки під кожен не потрібні.")
        rows.append([BASE + p, page_state(p), len(cs), sum(c["total"] for c in cs), ", ".join(c["id"] for c in cs),
                     f"{topl}: {topc}/{tot_n} ({topc / tot_n:.0%})", ver, why + " " + PAGE_NOTES.get(p, "")])
    parents = collections.OrderedDict()           # existing regional pages without own keywords, parents of matrix pages
    for p in pages:
        parts = p.split("/")
        if len(parts) == 5 and "/".join(parts[:4]) not in pages:
            parents.setdefault("/".join(parts[:4]), []).append(p)
    for par, kids in parents.items():
        rows.append([BASE + par, page_state(par), 0, 0, "—", "—", "", "Батьківська сторінка для матриць " + " і ".join("/" + x.rsplit("/", 1)[1] for x in kids) + " (власних ключів немає). " + PAGE_NOTES.get(par, "")])
    sheet(ws6, ["сторінка", "стан", "кластерів", "сума частотності", "кластери", "домінуючий тип у ТОП", "перевірка спірних", "чому так"], rows,
          [46, 10, 11, 16, 50, 36, 80, 110], wrap_cols=(6, 7))
    vf = os.path.join(ROOT, "landing-verification.json")
    if os.path.exists(vf):
        vr = json.load(open(vf, encoding="utf-8"))
        wsv = wb.create_sheet("Перевірка спірних")
        rowsv = []
        for x in vr:
            a_, b_ = x["a"], x["b"]
            rowsv.append([x["id"], x["why"], a_.get("url"), a_.get("status"), "; ".join(a_.get("h1") or []), a_.get("words"), a_.get("own_text_words"),
                          b_.get("url"), b_.get("status"), "; ".join(b_.get("h1") or []), b_.get("words"), b_.get("own_text_words"), x.get("text_overlap_6gram")])
        sheet(wsv, ["пара", "що порівнюємо", "URL A", "HTTP A", "H1 A", "слів A", "власний текст A", "URL B", "HTTP B", "H1 B", "слів B", "власний текст B", "збіг тексту (6-грами)"],
              rowsv, [34, 34, 60, 8, 40, 9, 12, 60, 8, 40, 9, 12, 12])
    ws7 = wb.create_sheet("Перетини між кластерами")
    sheet(ws7, ["кластер A", "кластер B", "макс. спільних URL", "ключі з макс. перетином", "висновок"], cross, [46, 46, 14, 70, 110], wrap_cols=(4,))
    ws8 = wb.create_sheet("SERP hard-4 (довідка)")
    sheet(ws8, ["кластер", "головний ключ", "сумарна частотність", "к-ть ключів", "ключі"],
          [[f"H{i:02d}", c[0], sum(K[k]["volume"] for k in c), len(c), "; ".join(c)] for i, c in
           enumerate(sorted(hard, key=lambda c: -sum(K[k]["volume"] for k in c)), 1)], [8, 46, 14, 10, 150])
    ws9 = wb.create_sheet("SERP soft-3 (довідка)")
    sheet(ws9, ["кластер", "головний ключ", "сумарна частотність", "к-ть ключів", "ключі"],
          [[f"S{i:02d}", c[0], sum(K[k]["volume"] for k in c), len(c), "; ".join(c)] for i, c in
           enumerate(sorted(soft, key=lambda c: -sum(K[k]["volume"] for k in c)), 1)], [8, 46, 14, 10, 150])
    try:
        wb.save(os.path.join(ROOT, "semantics-turkei.xlsx"))
    except PermissionError:   # file is open in Excel
        alt = os.path.join(ROOT, "semantics-turkei.new.xlsx")
        wb.save(alt)
        print("УВАГА: semantics-turkei.xlsx відкритий в Excel, збережено як", alt)

    # ---- cluster-map.html ----
    COL = ["#4E79A7", "#F28E2B", "#E15759", "#76B7B2", "#59A14F", "#EDC948", "#B07AA1", "#FF9DA7"]
    fam = [("Ця сторінка: /tour/turkei", lambda p: p == "/tour/turkei"), ("Pauschalreise", lambda p: p.endswith("/pauschalreise")),
           ("Last minute", lambda p: p.endswith("/last-minute")), ("Rundreisen", lambda p: p.endswith("/rundreisen")),
           ("Reisehinweise", lambda p: p.endswith("/reisehinweise")), ("All inclusive", lambda p: p.endswith("/all-inclusive")),
           ("Регіони (Side, Antalya, матриці)", lambda p: p.startswith("/tour/turkei/")), ("Готелі", lambda p: p.startswith("/hotels"))]
    cl = [{"name": n, "color": COL[i], "posts": []} for i, (n, _) in enumerate(fam)]
    for c in H:
        g = next(i for i, (_, f) in enumerate(fam) if f(c["page"]))
        cl[g]["posts"].append({"title": c["head"], "keyword": f"{c['id']} · {c['name']} · {len(c['members'])} ключ. → {c['page']}", "volume": c["total"],
                               "template": page_state(c["page"]), "wordCount": c["total"], "url": c["page"], "status": "planned"})
    keep = [i for i, c in enumerate(cl) if c["posts"]]
    links = []
    for new, old in enumerate(keep):
        for pi in range(len(cl[old]["posts"])):
            nid = f"cluster-{new}-post-{pi}"
            links.append({"from": nid, "to": "pillar", "type": "mandatory"})
            links.append({"from": "pillar", "to": nid, "type": "mandatory" if old == 0 else "recommended"})
    tot = sum(c["total"] for c in H)
    data = {"pillar": {"title": "/tour/turkei", "keyword": "türkei urlaub", "volume": K["türkei urlaub"]["volume"], "template": "поточна сторінка", "wordCount": tot, "url": "/tour/turkei"},
            "clusters": [cl[i] for i in keep], "links": links,
            "meta": {"totalPosts": len(H), "totalClusters": len(keep), "totalLinks": len(links), "estimatedWords": tot}}
    h = open(TEMPLATE, encoding="utf-8").read()
    s, e = h.index("const CLUSTER_DATA = {"), h.index("// === END CLUSTER DATA ===")
    h = h[:s] + "const CLUSTER_DATA = " + json.dumps(data, ensure_ascii=False) + ";\n    " + h[e:]
    rep = {'lang="en"': 'lang="uk"', "Content Cluster Map": "Карта кластерів: Reisen in die Türkei", '"Cluster map for: "': '"Карта кластерів для: "',
           "Keyword:": "Кластер:", "Volume:": "Сума частотності:", "Template:": "Стан сторінки:", "Words:": "Σ частотність:", "Status:": "Статус:",
           "Est. Words": "Σ частотність", "Mandatory link": "Лінк (ця сторінка ↔ кластер)", "Recommended link": "Рекомендований лінк (хаб → окрема сторінка)",
           "Optional link": "Додатковий лінк", '"Pillar page: "': '"Сторінка: "', '"Spoke page: "': '"Кластер: "', "Total Posts": "Кластерів",
           '<div class="stat-label">Clusters</div>': '<div class="stat-label">Груп (сторінок)</div>', "Internal Links": "Лінків",
           "</div> Written": "</div> Написано", "</div> Planned": "</div> Рекомендовано",
           "Generated by Claude SEO": "Створено Claude SEO · сторінки за типом посадкових конкурентів"}
    for k, v in rep.items():
        h = h.replace(k, v)
    open(os.path.join(ROOT, "cluster-map.html"), "w", encoding="utf-8").write(h)
    json.dump({"clusters": H, "cross": cross,
               "decisions": {k: {x: y for x, y in d.items()} for k, d in dec.items()},
               "types": {k: dict(v) for k, v in kw_types.items()}},
              open(os.path.join(ROOT, "clusters.json"), "w", encoding="utf-8"), ensure_ascii=False, indent=1)

    # ---- console summary ----
    print(f"keywords {len(K)}; clusters {len(H)} (lemma clusters {len(base)}); pairs with 0 shared URLs {dist[0] / sum(dist.values()):.1%}; SERP hard-4 {len(hard)}, soft-3 {len(soft)}")
    print("types:", collections.Counter(r[4] for r in URLS).most_common())
    for p, cs in pages.items():
        print(f"{p:<42} {page_state(p):<7} clusters {','.join(c['id'] for c in cs)} vol {sum(c['total'] for c in cs)}")
    print("statuses:", collections.Counter(d["status"] for d in dec.values()))


if __name__ == "__main__":
    main()
